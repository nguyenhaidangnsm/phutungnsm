# -*- coding: utf-8 -*-
"""
CHỈ TIÊU HMS theo từng cửa hàng  (thay file Excel "THEO DÕI CHỈ TIÊU HMS").

Luồng:
  * Cửa hàng : tải 4 file xuất từ HMS (xuất bán lẻ, xuất dịch vụ, nhận hàng, đặt hàng) -> xem kết quả của CHÍNH MÌNH.
  * Admin    : nhập bảng tra dùng chung (Part Category / Hao Mòn / Oil), chỉ tiêu gốc + tỉ lệ kế hoạch cho TỪNG cửa hàng,
               xem bảng tổng hợp mọi cửa hàng, bấm 1 dòng để xem chi tiết như cửa hàng đó.
  * "Realtime": mọi thay đổi tăng 1 số phiên bản (ct_meta). Trình duyệt hỏi /api/chi-tieu/version mỗi vài giây
               (1 câu SELECT nhỏ) và tự tải lại khi số phiên bản đổi.

Bảng tự tạo ở lần dùng đầu tiên (không cần sửa init_db).
Tính toán chạy ở server bằng pandas; kết quả được cache trong RAM theo (cửa hàng, tháng, phiên bản).
"""
import io
import re
import time
import threading
import unicodedata
from datetime import date, datetime
import calendar

import numpy as np
import pandas as pd
from flask import Blueprint, request, jsonify, session, send_file
from psycopg2.extras import execute_values

from app import get_db, vn_now, _valid_store_codes

chi_tieu_bp = Blueprint('chi_tieu', __name__)

# ----------------------------------------------------------------------------
# Hằng số theo đúng file Excel
# ----------------------------------------------------------------------------
KINDS = ('otc', 'jc', 'nhan', 'dat')
KIND_LABEL = {'otc': 'Xuất bán lẻ', 'jc': 'Xuất dịch vụ', 'nhan': 'Nhận hàng', 'dat': 'Đặt hàng'}
# Từ khoá tên sheet trong file Excel gốc (để tải thẳng cả file Excel cũng nhận đúng sheet)
KIND_SHEET_HINT = {'otc': 'xuat ban le', 'jc': 'xuat dich vu', 'nhan': 'nhan hang', 'dat': 'dat hang'}

PT_CATES = ('BP', 'GR', 'PM')                      # "Phụ tùng"
HM_G1 = ('Dây đai', 'Lọc gió', 'Má phanh', 'Nhông xích')
HM_G2 = ('Lốp', 'Bình điện', 'Nước làm mát', 'Bugi')
HM_ALL = HM_G1 + HM_G2
DEFAULT_PCT_XUAT = 1.02077
DEFAULT_PCT_DAT = 0.9308
DEFAULT_PCT_NHAN = 0.9308
NHAN_BAN_LOW, NHAN_BAN_HIGH = 0.95, 1.05

_ensure_lock = threading.Lock()
_tables_ready = False


def _json_err(msg, code=400):
    return jsonify({'error': msg}), code


def _strip_accents(s):
    s = unicodedata.normalize('NFD', str(s))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return s.replace('đ', 'd').replace('Đ', 'D')


def _key(s):
    """Chuẩn hoá tiêu đề cột/tên sheet: bỏ dấu, chữ thường, bỏ ký tự lạ."""
    return re.sub(r'[^a-z0-9]+', ' ', _strip_accents(s).lower()).strip()


# ----------------------------------------------------------------------------
# Tạo bảng
# ----------------------------------------------------------------------------
def ensure_tables():
    global _tables_ready
    if _tables_ready:
        return
    with _ensure_lock:
        if _tables_ready:
            return
        db = get_db()
        cur = db.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS ct_rows (
                id BIGSERIAL PRIMARY KEY,
                store_code TEXT NOT NULL,
                month CHAR(7) NOT NULL,          -- 'YYYY-MM'
                kind TEXT NOT NULL,              -- otc | jc | nhan | dat
                part_code TEXT NOT NULL,
                qty DOUBLE PRECISION NOT NULL DEFAULT 0,
                price DOUBLE PRECISION NOT NULL DEFAULT 0,   -- giá nhập (DNP) đã nhân 1000 theo quy tắc Excel
                dt DATE                          -- ngày bán/nhận (đã sửa lỗi đảo ngày-tháng); NULL với đặt hàng
            );
            CREATE INDEX IF NOT EXISTS ct_rows_sm ON ct_rows (store_code, month, kind);
            CREATE TABLE IF NOT EXISTS ct_uploads (
                store_code TEXT NOT NULL, month CHAR(7) NOT NULL, kind TEXT NOT NULL,
                filename TEXT, n_rows INT, n_skipped INT, uploaded_by TEXT, uploaded_at TIMESTAMP,
                PRIMARY KEY (store_code, month, kind)
            );
            CREATE TABLE IF NOT EXISTS ct_targets (
                store_code TEXT NOT NULL, month CHAR(7) NOT NULL,
                visits DOUBLE PRECISION DEFAULT 0,
                pt DOUBLE PRECISION DEFAULT 0, oil DOUBLE PRECISION DEFAULT 0, pg DOUBLE PRECISION DEFAULT 0,
                hm1 DOUBLE PRECISION DEFAULT 0, hm2 DOUBLE PRECISION DEFAULT 0,
                pct_xuat DOUBLE PRECISION DEFAULT 1.02077,
                pct_dat DOUBLE PRECISION DEFAULT 0.9308,
                pct_nhan DOUBLE PRECISION DEFAULT 0.9308,
                updated_by TEXT, updated_at TIMESTAMP,
                PRIMARY KEY (store_code, month)
            );
            CREATE TABLE IF NOT EXISTS ct_stores (
                store_code TEXT PRIMARY KEY, head_code TEXT, head_name TEXT
            );
            CREATE TABLE IF NOT EXISTS ct_ref_part (code TEXT PRIMARY KEY, ptype TEXT, cate TEXT);
            CREATE TABLE IF NOT EXISTS ct_ref_hm   (code TEXT PRIMARY KEY, grp TEXT);
            CREATE TABLE IF NOT EXISTS ct_ref_oil  (code TEXT PRIMARY KEY, kind TEXT, vehicle TEXT);
            CREATE TABLE IF NOT EXISTS ct_meta (key TEXT PRIMARY KEY, ver BIGINT NOT NULL);
        """)
        db.commit()
        cur.close()
        _tables_ready = True


# ----------------------------------------------------------------------------
# Phiên bản (realtime) + cache
# ----------------------------------------------------------------------------
def _bump(cur, *keys):
    ver = int(time.time() * 1000)
    for k in keys:
        cur.execute("""INSERT INTO ct_meta (key, ver) VALUES (%s, %s)
                       ON CONFLICT (key) DO UPDATE SET ver = GREATEST(ct_meta.ver + 1, EXCLUDED.ver)""", (k, ver))


def _skey(store, month):
    return f's:{store}:{month}'


_cache = {}              # (store, month) -> (ver_tuple, result)
_cache_lock = threading.Lock()
_ref_cache = {'ver': None, 'part': {}, 'hm': {}, 'oil': {}}


def _get_vers(cur, store, month):
    cur.execute("SELECT key, ver FROM ct_meta WHERE key IN ('ref', %s)", (_skey(store, month),))
    d = {r['key']: r['ver'] for r in cur.fetchall()}
    return d.get('ref', 0), d.get(_skey(store, month), 0)


def _load_ref(cur, ref_ver):
    """Bảng tra dùng chung, nạp vào RAM (~60k dòng) và giữ tới khi admin nạp lại."""
    if _ref_cache['ver'] == ref_ver and _ref_cache['part']:
        return _ref_cache
    with _cache_lock:
        cur.execute("SELECT code, ptype, cate FROM ct_ref_part")
        part = {r['code']: (r['ptype'] or '', (r['cate'] or '').upper()) for r in cur.fetchall()}
        cur.execute("SELECT code, grp FROM ct_ref_hm")
        hm = {r['code']: r['grp'] for r in cur.fetchall()}
        cur.execute("SELECT code, kind, vehicle FROM ct_ref_oil")
        oil = {r['code']: (r['kind'] or '', r['vehicle'] or '') for r in cur.fetchall()}
        _ref_cache.update(ver=ref_ver, part=part, hm=hm, oil=oil)
        _cache.clear()
    return _ref_cache


# ----------------------------------------------------------------------------
# Phân quyền
# ----------------------------------------------------------------------------
def _who():
    """-> (role, store_code). Admin đang 'mượn quyền' cửa hàng thì session['role'] đã là 'store'."""
    if 'user' not in session:
        return None, None
    return session.get('role'), session.get('store_code')


def _resolve_store(cur, role, own_store, requested):
    """Cửa hàng chỉ xem được của mình; admin xem được mọi cửa hàng hợp lệ."""
    if role == 'store':
        return own_store
    if role == 'admin':
        if requested and requested in _valid_store_codes(cur):
            return requested
        return None
    return None


def _month_arg():
    m = (request.values.get('month') or '').strip()
    if re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', m):
        return m
    return vn_now().strftime('%Y-%m')


# ----------------------------------------------------------------------------
# Đọc file Excel/CSV
# ----------------------------------------------------------------------------
def _norm_code(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r'\s+', '', str(v)).upper()


def _read_csv_loose(raw):
    """Đọc CSV 'lỏng': file HMS có vài dòng đầu (tên báo cáo, bộ lọc...) ít cột hơn dòng dữ liệu,
    nên pd.read_csv mặc định báo 'Expected 1 fields... saw N'. Dùng csv.reader + đệm cột cho đều."""
    import csv
    text = None
    for enc in ('utf-8-sig', 'utf-16', 'cp1258', 'latin-1'):
        try:
            text = raw.decode(enc)
            break
        except Exception:
            continue
    if text is None:
        text = raw.decode('latin-1', errors='ignore')
    lines = text.splitlines()
    # dò dấu phân cách theo dòng có nhiều cột nhất trong 30 dòng đầu
    best, best_n = ',', 0
    for d in (',', ';', '\t', '|'):
        n = max((len(next(csv.reader([ln], delimiter=d), [])) for ln in lines[:30] if ln.strip()), default=0)
        if n > best_n:
            best, best_n = d, n
    rows = [r for r in csv.reader(lines, delimiter=best)]
    width = max((len(r) for r in rows), default=0)
    rows = [[(c.strip() if c.strip() != '' else None) for c in r] + [None] * (width - len(r)) for r in rows]
    return pd.DataFrame(rows, dtype=object)


def _read_sheet(file_storage, sheet_hint=None, max_header_scan=12, header_keys=None):
    """Đọc 1 sheet; tự dò dòng tiêu đề (file HMS có dòng tiêu đề ở vị trí khác nhau).
    header_keys: tập tên cột (đã _key) - dòng nào chứa >=1 tên này được coi là dòng tiêu đề."""
    raw = file_storage.read()
    name = (file_storage.filename or '').lower()
    bio = io.BytesIO(raw)
    if name.endswith('.csv'):
        df = _read_csv_loose(raw)
    else:
        xl = pd.ExcelFile(bio)
        sheet = xl.sheet_names[0]
        if sheet_hint and len(xl.sheet_names) > 1:
            for want_prefix in (True, False):      # ưu tiên sheet "N. Nhập dữ liệu ..." rồi mới tới sheet trùng tên khác
                found = next((s for s in xl.sheet_names
                              if sheet_hint in _key(s) and (not want_prefix or 'nhap du lieu' in _key(s))), None)
                if found:
                    sheet = found
                    break
        df = xl.parse(sheet, header=None, dtype=object)
    hdr = None
    for i in range(min(max_header_scan, len(df))):
        keys = {_key(x) for x in df.iloc[i].tolist() if x is not None and not (isinstance(x, float) and np.isnan(x))}
        if header_keys and keys & set(header_keys):
            hdr = i
            break
    if hdr is None:
        raise ValueError('Không tìm thấy dòng tiêu đề trong file (kiểm tra lại đúng file xuất từ HMS).')
    cols = []
    seen = {}
    for x in df.iloc[hdr].tolist():
        k = _key(x) if x is not None and not (isinstance(x, float) and np.isnan(x)) else ''
        if k in seen:
            seen[k] += 1
            k = f'{k}__{seen[k]}'
        else:
            seen[k] = 0
        cols.append(k)
    out = df.iloc[hdr + 1:].copy()
    out.columns = cols
    return out.reset_index(drop=True)


def _pick(df, *cands):
    for c in cands:
        if c in df.columns:
            return c
    return None


def _parse_price(v, times1000=True):
    """Quy tắc Excel: nếu là chữ có dấu phẩy -> bỏ dấu phẩy (đã là VND đầy đủ); ngược lại nhân 1000."""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return 0.0
    if isinstance(v, str):
        s = re.sub(r'[^\d.,\-]', '', v.strip())          # bỏ ký hiệu tiền tệ (vd "₫2,577,750"), khoảng trắng
        if not s:
            return 0.0
        if ',' in s:                                      # có dấu phẩy: đã là VND đầy đủ, không nhân 1000
            try:
                return float(s.replace(',', ''))
            except ValueError:
                return 0.0
        try:
            x = float(s)
        except ValueError:
            return 0.0
        return x * 1000.0 if times1000 else x
    try:
        x = float(v)
    except (TypeError, ValueError):
        return 0.0
    return x * 1000.0 if times1000 else x


def _parse_dates(series, month):
    """Trả về Series ngày (date) đã sửa lỗi đảo ngày/tháng.
    HMS xuất dd/mm/yyyy nhưng khi Excel hiểu theo kiểu Mỹ thì 04/10 thành 10/04 (ngày 10 tháng 4).
    Nếu ngày đọc được KHÔNG thuộc tháng đang xét mà số 'ngày' trùng số tháng đang xét -> đảo lại."""
    y, mth = int(month[:4]), int(month[5:7])
    res = []
    for v in series.tolist():
        d = None
        if v is None or (isinstance(v, float) and np.isnan(v)):
            res.append(None)
            continue
        if isinstance(v, (datetime, pd.Timestamp)):
            d = v
            was_dt = True
        elif isinstance(v, date):
            d = datetime(v.year, v.month, v.day)
            was_dt = True
        elif isinstance(v, (int, float)) and 20000 < float(v) < 80000:      # số serial Excel
            d = datetime(1899, 12, 30) + pd.to_timedelta(float(v), unit='D')
            was_dt = True
        else:
            s = str(v).strip().split(' ')[0]
            d = pd.to_datetime(s, dayfirst=True, errors='coerce')
            was_dt = False
        if d is None or pd.isna(d):
            res.append(None)
            continue
        if was_dt and not (d.year == y and d.month == mth) and d.day == mth and d.month <= 12 and d.year == y:
            try:
                d = datetime(d.year, d.day, d.month)
            except ValueError:
                pass
        res.append(date(d.year, d.month, d.day))
    return pd.Series(res, index=series.index, dtype=object)


# Cột cần cho từng loại file (đã _key)
SPEC = {
    'otc':  dict(code=('ma phu tung',), qty=('so luong xac nhan',), price=('gia nhap',),
                 date=('ngay hoa don', 'ngay xuat', 'ngay gio duoc cap nhat'), times1000=True),
    'jc':   dict(code=('ma phu tung',), qty=('so luong xac nhan',), price=('gia nhap',),
                 date=('ngay gio duoc cap nhat', 'ngay hoa don', 'ngay xuat'), times1000=True),
    'nhan': dict(code=('part',), qty=('qty',), price=('dnp unit price', 'dnp'),
                 date=('mrn date', 'transaction date'), times1000=True),
    'dat':  dict(code=('part number',), qty=('order quantity',), price=('dnp',),
                 date=('po date',), times1000=False),
}


def _parse_kind_file(kind, file_storage, month):
    spec = SPEC[kind]
    df = _read_sheet(file_storage, KIND_SHEET_HINT[kind], header_keys=set(spec['code']))
    c_code, c_qty, c_price = _pick(df, *spec['code']), _pick(df, *spec['qty']), _pick(df, *spec['price'])
    c_date = _pick(df, *spec['date'])
    if kind == 'nhan':
        # File MRN của HMS có cả phiếu "Open" (chưa đóng, hàng chưa tính là đã nhận) -> chỉ lấy "Closed".
        # File không có cột MRN Status (định dạng cũ) thì giữ nguyên, không lọc.
        c_mrn = _pick(df, 'mrn status')
        if c_mrn is not None:
            df = df[df[c_mrn].map(lambda v: _key(str(v)) == 'closed')]
    if kind == 'dat':
        # File đặt hàng (PO): tính mọi dòng, trừ phiếu có PO Status = Cancelled (đã huỷ). Không có cột này thì giữ nguyên.
        c_po = _pick(df, 'po status')
        if c_po is not None:
            df = df[~df[c_po].map(lambda v: _key(str(v)).startswith('cancel'))]
    miss = [n for n, c in (('Mã phụ tùng', c_code), ('Số lượng', c_qty), ('Giá nhập', c_price)) if c is None]
    if miss:
        raise ValueError('File thiếu cột: ' + ', '.join(miss) + '. Hãy chọn đúng file "' + KIND_LABEL[kind] + '" xuất từ HMS.')
    out = pd.DataFrame({
        'part_code': df[c_code].map(_norm_code),
        'qty': pd.to_numeric(df[c_qty], errors='coerce').fillna(0.0),
        'price': df[c_price].map(lambda v: _parse_price(v, spec['times1000'])),
    })
    skipped = 0
    if kind == 'dat':
        out['dt'] = None
    else:
        if c_date is None:
            raise ValueError('File thiếu cột ngày. Hãy chọn đúng file "' + KIND_LABEL[kind] + '".')
        out['dt'] = _parse_dates(df[c_date], month)
    total_before = len(out)
    out = out[(out['part_code'] != '') & (out['part_code'] != '0')]
    if kind != 'dat':
        in_month = out['dt'].map(lambda d: d is not None and f'{d.year:04d}-{d.month:02d}' == month)
        skipped = int((~in_month).sum())
        out = out[in_month]
    return out.reset_index(drop=True), skipped


# ----------------------------------------------------------------------------
# Tính toán
# ----------------------------------------------------------------------------
def _week_of(d):
    """Tuần theo thứ Hai-Chủ nhật, tuần 1 là tuần chứa ngày mùng 1 (Chủ nhật 04/10/2026 thuộc Tuần 1)."""
    return (d.day - 1 + date(d.year, d.month, 1).weekday()) // 7 + 1


def _prepare(df, ref, with_date=True):
    if df.empty:
        df = df.assign(cate='', hm='', grp='', amount=0.0)
        return df
    part, hm, oil = ref['part'], ref['hm'], ref['oil']
    codes = df['part_code']
    df = df.copy()
    df['cate'] = codes.map(lambda c: part.get(c, ('', 'OTHER'))[1] or 'OTHER')
    df['grp'] = codes.map(lambda c: hm.get(c, ''))
    df['hm'] = df['grp'] != ''
    df['amount'] = df['qty'] * df['price']
    return df


def _sum(df, mask=None):
    if df.empty:
        return 0.0
    return float((df['amount'] if mask is None else df.loc[mask, 'amount']).sum())


def _by_cate(df):
    r = {c: 0.0 for c in ('BP', 'GR', 'PM', 'OIL', 'PK', 'PG', 'OTHER')}
    if not df.empty:
        for c, v in df.groupby('cate')['amount'].sum().items():
            r[c if c in r else 'OTHER'] += float(v)
    return r


def _by_grp(df):
    r = {g: 0.0 for g in HM_ALL}
    if not df.empty:
        for g, v in df[df['grp'] != ''].groupby('grp')['amount'].sum().items():
            if g in r:
                r[g] += float(v)
    return r


def _compute(cur, store, month, ref):
    cur.execute("SELECT * FROM ct_targets WHERE store_code=%s AND month=%s", (store, month))
    t = cur.fetchone()
    tg = dict(visits=0, pt=0, oil=0, pg=0, hm1=0, hm2=0,
              pct_xuat=DEFAULT_PCT_XUAT, pct_dat=DEFAULT_PCT_DAT, pct_nhan=DEFAULT_PCT_NHAN)
    has_target = False
    if t:
        has_target = True
        for k in tg:
            if t.get(k) is not None:
                tg[k] = float(t[k])
    cur.execute("SELECT kind, part_code, qty, price, dt FROM ct_rows WHERE store_code=%s AND month=%s", (store, month))
    rows = cur.fetchall()
    base = pd.DataFrame(rows, columns=['kind', 'part_code', 'qty', 'price', 'dt']) if rows else \
        pd.DataFrame(columns=['kind', 'part_code', 'qty', 'price', 'dt'])
    frames = {k: _prepare(base[base['kind'] == k].drop(columns='kind'), ref) for k in KINDS}
    otc, jc, nhan, dat = frames['otc'], frames['jc'], frames['nhan'], frames['dat']

    c_otc, c_jc, c_nhan, c_dat = _by_cate(otc), _by_cate(jc), _by_cate(nhan), _by_cate(dat)

    def pt(c):  return c['BP'] + c['GR'] + c['PM']
    # Xuất = Phụ tùng + Dầu nhớt + Phụ gia (đúng F14 trong Excel). Đặt/Nhận tính cùng phạm vi để tỉ lệ so sánh được.
    def tot(c): return pt(c) + c['OIL'] + c['PG']

    target_total = tg['pt'] + tg['oil'] + tg['pg']
    xuat_jc, xuat_otc = tot(c_jc), tot(c_otc)
    xuat = xuat_jc + xuat_otc
    dat_v, nhan_v = tot(c_dat), tot(c_nhan)
    plan_xuat = tg['pct_xuat'] * target_total
    plan_dat = tg['pct_dat'] * plan_xuat
    plan_nhan = tg['pct_nhan'] * plan_xuat
    ratio = (nhan_v / xuat) if xuat else None

    # Hao mòn theo nhóm
    g_dat, g_nhan = _by_grp(dat), _by_grp(nhan)
    g_otc, g_jc = _by_grp(otc), _by_grp(jc)
    g_ban = {g: g_otc[g] + g_jc[g] for g in HM_ALL}

    def grp_row(groups, target):
        n, b, d = sum(g_nhan[g] for g in groups), sum(g_ban[g] for g in groups), sum(g_dat[g] for g in groups)
        return dict(nhan=n, xuat=b, dat=d, target=target,
                    ratio=(n / b) if b else None, done=(b / target) if target else None)
    hm1, hm2 = grp_row(HM_G1, tg['hm1']), grp_row(HM_G2, tg['hm2'])

    # Chi tiết theo nhóm tháng
    def hm_total(df): return _sum(df, df['hm']) if not df.empty else 0.0
    chi_tiet = {
        'cols': ['BP', 'GR', 'PM', 'OIL', 'PK', 'PG', 'HM'],
        'ban': [c_otc[c] + c_jc[c] for c in ('BP', 'GR', 'PM', 'OIL', 'PK', 'PG')] + [hm_total(otc) + hm_total(jc)],
        'dat': [c_dat[c] for c in ('BP', 'GR', 'PM', 'OIL', 'PK', 'PG')] + [hm_total(dat)],
        'nhan': [c_nhan[c] for c in ('BP', 'GR', 'PM', 'OIL', 'PK', 'PG')] + [hm_total(nhan)],
        'sua_chua': xuat_jc, 'ban_le': xuat_otc,
    }
    pt_bl, pt_sc = pt(c_otc), pt(c_jc)
    ban_le_sc = {
        'ban_le': dict(pt=pt_bl, oil=c_otc['OIL'], hm=hm_total(otc)),
        'sua_chua': dict(pt=pt_sc, oil=c_jc['OIL'], hm=hm_total(jc)),
    }

    # Nhớt dịch vụ (chỉ nhớt máy: phuy/chai x xe ga/xe số)
    oilmap = ref['oil']
    nhot = {'phuy': {'ga': [0.0, 0.0, 0], 'so': [0.0, 0.0, 0]}, 'chai': {'ga': [0.0, 0.0, 0], 'so': [0.0, 0.0, 0]}}   # [SL, tiền, số dòng]
    if not jc.empty:
        for r in jc[['part_code', 'qty', 'amount']].itertuples(index=False):
            o = oilmap.get(r.part_code)
            if not o:
                continue
            kind = 'phuy' if 'PHUY' in o[0].upper() else ('chai' if 'CHAI' in o[0].upper() else None)
            veh = 'ga' if 'GA' in _strip_accents(o[1]).upper() else ('so' if 'SO' in _strip_accents(o[1]).upper() else None)
            if kind and veh:
                nhot[kind][veh][0] += float(r.qty)
                nhot[kind][veh][1] += float(r.amount)
                nhot[kind][veh][2] += 1

    # Theo tuần / theo ngày (xuất)
    last_day = calendar.monthrange(int(month[:4]), int(month[5:7]))[1]
    first = date(int(month[:4]), int(month[5:7]), 1)
    n_weeks = _week_of(date(first.year, first.month, last_day))
    def add_week(df):
        if df.empty:
            return df.assign(wk=pd.Series(dtype=int))
        return df.assign(wk=df['dt'].map(_week_of))
    otc_w, jc_w, nhan_w = add_week(otc), add_week(jc), add_week(nhan)
    tuan_xuat = []
    for w in range(1, n_weeks + 1):
        o, j = otc_w[otc_w['wk'] == w] if not otc_w.empty else otc_w, jc_w[jc_w['wk'] == w] if not jc_w.empty else jc_w
        row = dict(tuan=w, ban_le=_sum(o), dich_vu=_sum(j),
                   hao_mon=(_sum(o, o['hm']) + _sum(j, j['hm'])))
        row['tong'] = row['ban_le'] + row['dich_vu']
        go, gj = _by_grp(o), _by_grp(j)
        row['nhom'] = {g: go[g] + gj[g] for g in HM_ALL}
        tuan_xuat.append(row)
    tuan_nhan = []
    for w in range(1, n_weeks + 1):
        n = nhan_w[nhan_w['wk'] == w] if not nhan_w.empty else nhan_w
        c = _by_cate(n)
        row = dict(tuan=w, pt=pt(c), oil=c['OIL'], pg=c['PG'], hao_mon=(_sum(n, n['hm']) if not n.empty else 0.0))
        row['tong'] = row['pt'] + row['oil'] + row['pg']
        tuan_nhan.append(row)
    ngay = []
    for d in range(1, last_day + 1):
        dd = date(first.year, first.month, d)
        o = otc[otc['dt'] == dd] if not otc.empty else otc
        j = jc[jc['dt'] == dd] if not jc.empty else jc
        go, gj = _by_grp(o), _by_grp(j)
        ngay.append(dict(ngay=d, thu=dd.weekday(), ban_le=_sum(o), dich_vu=_sum(j),
                         nhom={g: [go[g], gj[g]] for g in HM_ALL}))

    other_amt = c_otc['OTHER'] + c_jc['OTHER'] + c_nhan['OTHER'] + c_dat['OTHER']
    other_codes = set()
    for df in (otc, jc, nhan, dat):
        if not df.empty:
            other_codes.update(df.loc[df['cate'] == 'OTHER', 'part_code'].tolist())

    cur.execute("SELECT kind, filename, n_rows, n_skipped, uploaded_by, uploaded_at FROM ct_uploads WHERE store_code=%s AND month=%s",
                (store, month))
    ups = {r['kind']: dict(file=r['filename'], rows=r['n_rows'], skipped=r['n_skipped'], by=r['uploaded_by'],
                           at=r['uploaded_at'].strftime('%d/%m/%Y %H:%M') if r['uploaded_at'] else '')
           for r in cur.fetchall()}

    return {
        'store': store, 'month': month, 'has_target': has_target,
        'target': dict(tg, total=target_total),
        'plan': dict(xuat=plan_xuat, dat=plan_dat, nhan=plan_nhan),
        'actual': dict(xuat=xuat, dat=dat_v, nhan=nhan_v, xuat_jc=xuat_jc, xuat_otc=xuat_otc),
        'gap': dict(xuat=plan_xuat - xuat, dat=plan_dat - dat_v, nhan=plan_nhan - nhan_v),
        'ratio_nhan_ban': ratio,
        'ratio_ok': (ratio is not None and NHAN_BAN_LOW <= ratio <= NHAN_BAN_HIGH),
        'pct_done': (xuat / target_total) if target_total else None,
        'hm1': hm1, 'hm2': hm2,
        'chi_tiet': chi_tiet, 'ban_le_sc': ban_le_sc, 'nhot': nhot,
        'tuan_xuat': tuan_xuat, 'tuan_nhan': tuan_nhan, 'ngay': ngay,
        'groups': list(HM_ALL),
        'uploads': ups,
        'warn': dict(other_amount=other_amt, other_codes=sorted(other_codes)[:30], other_count=len(other_codes),
                     ref_empty=not ref['part']),
        'computed_at': vn_now().strftime('%H:%M:%S'),
    }


def _result(cur, store, month):
    ref_ver, s_ver = _get_vers(cur, store, month)
    ref = _load_ref(cur, ref_ver)
    k = (store, month)
    with _cache_lock:
        hit = _cache.get(k)
    if hit and hit[0] == (ref_ver, s_ver):
        return hit[1]
    res = _compute(cur, store, month, ref)
    res['ver'] = f'{ref_ver}.{s_ver}'
    with _cache_lock:
        if len(_cache) > 400:
            _cache.clear()
        _cache[k] = ((ref_ver, s_ver), res)
    return res


# Cửa hàng không theo dõi chỉ tiêu HMS (ẩn khỏi tổng hợp, danh sách chọn, nhập chỉ tiêu, xuất Excel tổng hợp)
_NO_TRACK = {'NSM1'}


def _track_stores(cur):
    import re as _re
    return sorted((s for s in _valid_store_codes(cur) if s not in _NO_TRACK),
                  key=lambda s: [int(x) if x.isdigit() else x for x in _re.split(r'(\d+)', s)])


def _summary_row(res):
    t, a, p = res['target'], res['actual'], res['plan']
    flag = []
    if res['ratio_nhan_ban'] is not None and not res['ratio_ok']:
        flag.append('Nhận/bán ngoài 95–105%')
    if not res['has_target']:
        flag.append('Chưa có chỉ tiêu')
    if res['warn']['other_count']:
        flag.append(f"{res['warn']['other_count']} mã chưa phân loại")
    return dict(store=res['store'], target=t['total'], plan=p['xuat'], xuat=a['xuat'], pct=res['pct_done'],
                dat=a['dat'], nhan=a['nhan'], plan_dat=p['dat'], plan_nhan=p['nhan'],
                ratio=res['ratio_nhan_ban'], ratio_ok=res['ratio_ok'],
                hm1_done=res['hm1']['done'], hm2_done=res['hm2']['done'],
                hm1_ratio=res['hm1']['ratio'], hm2_ratio=res['hm2']['ratio'],
                hm1_target=res['hm1']['target'], hm1_xuat=res['hm1']['xuat'], hm2_target=res['hm2']['target'], hm2_xuat=res['hm2']['xuat'],
                has_target=res['has_target'], flags=flag,
                uploaded=len(res['uploads']), ver=res['ver'])


# ----------------------------------------------------------------------------
# API
# ----------------------------------------------------------------------------
@chi_tieu_bp.before_request
def _guard():
    if request.endpoint and request.endpoint.startswith('chi_tieu.'):
        role, _ = _who()
        if role not in ('admin', 'store'):
            return _json_err('Chưa đăng nhập hoặc không có quyền.', 401)
        ensure_tables()


@chi_tieu_bp.route('/api/chi-tieu/version')
def api_version():
    """Nhẹ nhất có thể - trình duyệt hỏi định kỳ để biết có cần tải lại không."""
    role, own = _who()
    month = _month_arg()
    cur = get_db().cursor()
    if role == 'admin' and not request.args.get('store'):
        cur.execute("SELECT COALESCE(MAX(ver),0) AS v FROM ct_meta WHERE key='ref' OR key LIKE %s", (f's:%:{month}',))
    else:
        store = _resolve_store(cur, role, own, request.args.get('store'))
        if not store:
            return _json_err('Cửa hàng không hợp lệ.')
        cur.execute("SELECT COALESCE(MAX(ver),0) AS v FROM ct_meta WHERE key IN ('ref', %s)", (_skey(store, month),))
    v = cur.fetchone()['v']
    cur.close()
    return jsonify({'v': int(v)})


@chi_tieu_bp.route('/api/chi-tieu/me')
def api_me():
    role, own = _who()
    cur = get_db().cursor()
    stores = _track_stores(cur) if role == 'admin' else [own]
    cur.execute("SELECT store_code, head_code, head_name FROM ct_stores")
    heads = {r['store_code']: dict(head_code=r['head_code'] or '', head_name=r['head_name'] or '') for r in cur.fetchall()}
    cur.execute("SELECT (SELECT COUNT(*) FROM ct_ref_part) p, (SELECT COUNT(*) FROM ct_ref_hm) h, (SELECT COUNT(*) FROM ct_ref_oil) o")
    rc = cur.fetchone()
    cur.close()
    return jsonify({'role': role, 'store': own, 'stores': stores, 'heads': heads,
                    'ref': {'part': rc['p'], 'hm': rc['h'], 'oil': rc['o']},
                    'month': vn_now().strftime('%Y-%m')})


@chi_tieu_bp.route('/api/chi-tieu/data')
def api_data():
    role, own = _who()
    month = _month_arg()
    cur = get_db().cursor()
    store = _resolve_store(cur, role, own, request.args.get('store'))
    if not store:
        return _json_err('Vui lòng chọn cửa hàng.')
    res = _result(cur, store, month)
    cur.close()
    return jsonify(res)


@chi_tieu_bp.route('/api/chi-tieu/overview')
def api_overview():
    role, _ = _who()
    if role != 'admin':
        return _json_err('Chỉ admin xem được tổng hợp.', 403)
    month = _month_arg()
    cur = get_db().cursor()
    rows = [_summary_row(_result(cur, s, month)) for s in _track_stores(cur)]
    cur.close()
    done = [r for r in rows if r['pct'] is not None and r['has_target']]
    done.sort(key=lambda r: -r['pct'])
    for i, r in enumerate(done, 1):
        r['rank'] = i
    tot_t = sum(r['target'] for r in rows)
    tot_x = sum(r['xuat'] for r in rows)
    return jsonify({'month': month, 'rows': rows,
                    'total': dict(target=tot_t, xuat=tot_x, pct=(tot_x / tot_t if tot_t else None),
                                  dat=sum(r['dat'] for r in rows), nhan=sum(r['nhan'] for r in rows))})


@chi_tieu_bp.route('/api/chi-tieu/upload', methods=['POST'])
def api_upload():
    """Cửa hàng (hoặc admin chọn cửa hàng) tải 1 hoặc nhiều trong 4 file. Mỗi file THAY THẾ dữ liệu cùng loại của tháng đó."""
    role, own = _who()
    month = _month_arg()
    db = get_db()
    cur = db.cursor()
    store = _resolve_store(cur, role, own, request.form.get('store'))
    if not store:
        return _json_err('Vui lòng chọn cửa hàng.')
    done, errors = [], []
    actor = session.get('full_name') or session.get('user')
    for kind in KINDS:
        f = request.files.get(kind)
        if not f or not f.filename:
            continue
        try:
            df, skipped = _parse_kind_file(kind, f, month)
        except Exception as e:
            errors.append(f'{KIND_LABEL[kind]}: {e}')
            continue
        try:
            cur.execute("DELETE FROM ct_rows WHERE store_code=%s AND month=%s AND kind=%s", (store, month, kind))
            if len(df):
                vals = [(store, month, kind, r.part_code, float(r.qty), float(r.price), r.dt)
                        for r in df.itertuples(index=False)]
                execute_values(cur, "INSERT INTO ct_rows (store_code, month, kind, part_code, qty, price, dt) VALUES %s",
                               vals, page_size=2000)
            cur.execute("""INSERT INTO ct_uploads (store_code, month, kind, filename, n_rows, n_skipped, uploaded_by, uploaded_at)
                           VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
                           ON CONFLICT (store_code, month, kind) DO UPDATE SET filename=EXCLUDED.filename,
                           n_rows=EXCLUDED.n_rows, n_skipped=EXCLUDED.n_skipped, uploaded_by=EXCLUDED.uploaded_by,
                           uploaded_at=EXCLUDED.uploaded_at""",
                        (store, month, kind, f.filename, len(df), skipped, actor, vn_now()))
            _bump(cur, _skey(store, month))
            db.commit()
            done.append(f'{KIND_LABEL[kind]}: {len(df)} dòng' + (f' (bỏ {skipped} dòng ngoài tháng {month[5:]}/{month[:4]})' if skipped else ''))
        except Exception as e:
            db.rollback()
            errors.append(f'{KIND_LABEL[kind]}: lỗi lưu dữ liệu ({e})')
    cur.close()
    if not done and not errors:
        return _json_err('Chưa chọn file nào.')
    return jsonify({'done': done, 'errors': errors})


@chi_tieu_bp.route('/api/chi-tieu/clear', methods=['POST'])
def api_clear():
    role, own = _who()
    month = _month_arg()
    db = get_db()
    cur = db.cursor()
    store = _resolve_store(cur, role, own, request.form.get('store'))
    kind = request.form.get('kind')
    if not store:
        return _json_err('Vui lòng chọn cửa hàng.')
    if kind not in KINDS:
        return _json_err('Loại dữ liệu không hợp lệ.')
    cur.execute("DELETE FROM ct_rows WHERE store_code=%s AND month=%s AND kind=%s", (store, month, kind))
    cur.execute("DELETE FROM ct_uploads WHERE store_code=%s AND month=%s AND kind=%s", (store, month, kind))
    _bump(cur, _skey(store, month))
    db.commit()
    cur.close()
    return jsonify({'ok': True})


# --- Admin: chỉ tiêu + tỉ lệ kế hoạch từng cửa hàng ---------------------------------------------
@chi_tieu_bp.route('/api/chi-tieu/targets')
def api_targets_get():
    role, _ = _who()
    if role != 'admin':
        return _json_err('Chỉ admin.', 403)
    month = _month_arg()
    cur = get_db().cursor()
    cur.execute("SELECT * FROM ct_targets WHERE month=%s", (month,))
    have = {r['store_code']: r for r in cur.fetchall()}
    out = []
    for s in _track_stores(cur):
        r = have.get(s)
        out.append(dict(store=s, has=bool(r),
                        visits=(r or {}).get('visits') or 0, pt=(r or {}).get('pt') or 0, oil=(r or {}).get('oil') or 0,
                        pg=(r or {}).get('pg') or 0, hm1=(r or {}).get('hm1') or 0, hm2=(r or {}).get('hm2') or 0,
                        pct_xuat=(r or {}).get('pct_xuat') or DEFAULT_PCT_XUAT,
                        pct_dat=(r or {}).get('pct_dat') or DEFAULT_PCT_DAT,
                        pct_nhan=(r or {}).get('pct_nhan') or DEFAULT_PCT_NHAN))
    cur.close()
    return jsonify({'month': month, 'rows': out})


def _num(v, default=0.0):
    try:
        return float(str(v).replace(',', '').strip())
    except (TypeError, ValueError):
        return default


@chi_tieu_bp.route('/api/chi-tieu/targets', methods=['POST'])
def api_targets_save():
    role, _ = _who()
    if role != 'admin':
        return _json_err('Chỉ admin.', 403)
    month = _month_arg()
    body = request.get_json(silent=True) or {}
    items = body.get('rows') or []
    db = get_db()
    cur = db.cursor()
    valid = _valid_store_codes(cur)
    actor = session.get('full_name') or session.get('user')
    n = 0
    for it in items:
        s = it.get('store')
        if s not in valid:
            continue
        px, pd_, pn = _num(it.get('pct_xuat'), DEFAULT_PCT_XUAT), _num(it.get('pct_dat'), DEFAULT_PCT_DAT), _num(it.get('pct_nhan'), DEFAULT_PCT_NHAN)
        # Cho phép nhập 102.077 (%) hoặc 1.02077
        px, pd_, pn = [x / 100.0 if x > 5 else x for x in (px, pd_, pn)]
        cur.execute("""INSERT INTO ct_targets (store_code, month, visits, pt, oil, pg, hm1, hm2, pct_xuat, pct_dat, pct_nhan, updated_by, updated_at)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                       ON CONFLICT (store_code, month) DO UPDATE SET visits=EXCLUDED.visits, pt=EXCLUDED.pt, oil=EXCLUDED.oil,
                       pg=EXCLUDED.pg, hm1=EXCLUDED.hm1, hm2=EXCLUDED.hm2, pct_xuat=EXCLUDED.pct_xuat, pct_dat=EXCLUDED.pct_dat,
                       pct_nhan=EXCLUDED.pct_nhan, updated_by=EXCLUDED.updated_by, updated_at=EXCLUDED.updated_at""",
                    (s, month, _num(it.get('visits')), _num(it.get('pt')), _num(it.get('oil')), _num(it.get('pg')),
                     _num(it.get('hm1')), _num(it.get('hm2')), px, pd_, pn, actor, vn_now()))
        _bump(cur, _skey(s, month))
        n += 1
    db.commit()
    cur.close()
    return jsonify({'saved': n})


def _blank(v):
    return v is None or (isinstance(v, float) and np.isnan(v))


def _find_row(df, label_part, start=0):
    """Tìm dòng có nhãn ở cột A hoặc B (file Excel của Honda khi thì để nhãn ở cột A, khi thì cột B).
    Ưu tiên ô trùng khớp hoàn toàn, sau đó mới tới ô chứa nhãn."""
    want = _key(label_part)
    for exact in (True, False):
        for i in range(start, len(df)):
            for ci in (0, 1):
                if df.shape[1] <= ci:
                    continue
                v = df.iloc[i, ci]
                if isinstance(v, str):
                    k = _key(v)
                    if (k == want) if exact else (want in k):
                        return i
    return None


def _parse_target_sheet(d):
    """Đọc 1 sheet chỉ tiêu của Honda -> dict(months={cột: 'YYYY-MM'}, head_code, head_name, rows={tên: dòng}).
    Trả về None nếu sheet không có dòng 'Tháng' (không phải sheet chỉ tiêu)."""
    hrow = _find_row(d, 'Tháng')
    if hrow is None:
        return None
    months = {}
    for ci in range(2, d.shape[1]):
        m = re.fullmatch(r'T(\d{1,2})\.(\d{4})', str(d.iloc[hrow, ci]).strip())
        if m:
            months[ci] = f'{m.group(2)}-{int(m.group(1)):02d}'
    if not months:
        return None
    rows = dict(
        vis=_find_row(d, 'Lượng khách', hrow), pt=_find_row(d, 'Phụ tùng', hrow), oil=_find_row(d, 'Dầu nhớt', hrow),
        pg=_find_row(d, 'Phụ gia', hrow), hm1=_find_row(d, 'Nhóm 1', hrow), hm2=_find_row(d, 'Nhóm 2', hrow))

    def first_val(label):
        r = _find_row(d, label)
        if r is None:
            return None
        for ci in range(2, d.shape[1]):
            v = d.iloc[r, ci]
            if not _blank(v) and str(v).strip():
                return v
        return None
    hc, hn = first_val('Mã HEAD'), first_val('Tên HEAD')
    return dict(d=d, months=months, rows=rows,
                head_code=_norm_code(hc) if hc is not None else None,
                head_name=str(hn).strip() if hn is not None else None)


def _save_target_sheet(cur, store, ps, actor):
    """Ghi chỉ tiêu gốc của 1 cửa hàng, mọi tháng trong sheet. Giữ nguyên tỉ lệ kế hoạch admin đã đặt."""
    d, rows = ps['d'], ps['rows']

    def val(r, ci):
        if r is None or _blank(d.iloc[r, ci]):
            return 0.0
        return _num(d.iloc[r, ci], 0.0)
    n = 0
    for ci, m in ps['months'].items():
        vis, pt_, oil, pg = (val(rows[k], ci) for k in ('vis', 'pt', 'oil', 'pg'))
        h1, h2 = val(rows['hm1'], ci), val(rows['hm2'], ci)
        if not any((vis, pt_, oil, pg, h1, h2)):
            continue
        cur.execute("""INSERT INTO ct_targets (store_code, month, visits, pt, oil, pg, hm1, hm2, updated_by, updated_at)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                       ON CONFLICT (store_code, month) DO UPDATE SET visits=EXCLUDED.visits, pt=EXCLUDED.pt, oil=EXCLUDED.oil,
                       pg=EXCLUDED.pg, hm1=EXCLUDED.hm1, hm2=EXCLUDED.hm2, updated_by=EXCLUDED.updated_by, updated_at=EXCLUDED.updated_at""",
                    (store, m, vis, pt_, oil, pg, h1, h2, actor, vn_now()))
        _bump(cur, _skey(store, m))
        n += 1
    if ps['head_code']:
        cur.execute("""INSERT INTO ct_stores (store_code, head_code, head_name) VALUES (%s,%s,%s)
                       ON CONFLICT (store_code) DO UPDATE SET head_code=EXCLUDED.head_code,
                       head_name=COALESCE(EXCLUDED.head_name, ct_stores.head_name)""",
                    (store, ps['head_code'], ps['head_name']))
    return n


def _store_for_sheet(cur, ps, valid, used):
    """Đoán cửa hàng của 1 sheet: (1) mã HEAD đã gắn với cửa hàng từ lần nạp trước,
    (2) tên HEAD dạng '... #3' -> NS3, (3) tên HEAD/sheet chứa đúng mã cửa hàng (vd 'NS3')."""
    if ps['head_code']:
        cur.execute("SELECT store_code FROM ct_stores WHERE head_code=%s", (ps['head_code'],))
        for r in cur.fetchall():
            if r['store_code'] in valid and r['store_code'] not in used:
                return r['store_code']
    name = ps['head_name'] or ''
    m = re.search(r'#\s*(\d+)\s*$', name)
    if m and f'NS{m.group(1)}' in valid and f'NS{m.group(1)}' not in used:
        return f'NS{m.group(1)}'
    for sc in sorted(valid, key=len, reverse=True):
        if re.search(r'(?<![A-Z0-9])' + re.escape(sc) + r'(?![A-Z0-9])', name.upper()) and sc not in used:
            return sc
    return None


@chi_tieu_bp.route('/api/chi-tieu/targets/import', methods=['POST'])
def api_targets_import():
    """Nạp chỉ tiêu GỐC từ file Excel của Honda (tất cả các tháng trong file).
    - store = '__all__' (hoặc bỏ trống): file có NHIỀU SHEET, mỗi sheet 1 cửa hàng -> tự nhận cửa hàng của từng sheet.
    - store = mã cửa hàng: nạp 1 cửa hàng (sheet 'Chỉ tiêu gốc' nếu có, không thì sheet đầu tiên)."""
    role, _ = _who()
    if role != 'admin':
        return _json_err('Chỉ admin.', 403)
    f = request.files.get('file')
    store = (request.form.get('store') or '__all__').strip()
    db = get_db()
    cur = db.cursor()
    valid = _valid_store_codes(cur)
    if store != '__all__' and store not in valid:
        return _json_err('Vui lòng chọn cửa hàng.')
    if not f:
        return _json_err('Chưa chọn file.')
    try:
        xl = pd.ExcelFile(io.BytesIO(f.read()))
        actor = session.get('full_name') or session.get('user')
        if store != '__all__':
            sheet = next((s for s in xl.sheet_names if 'chi tieu goc' in _key(s)), xl.sheet_names[0])
            ps = _parse_target_sheet(xl.parse(sheet, header=None, dtype=object))
            if ps is None:
                return _json_err('Không thấy dòng "Tháng" (T4.2026, T5.2026...) trong sheet đầu tiên của file.')
            n = _save_target_sheet(cur, store, ps, actor)
            db.commit()
            cur.close()
            return jsonify({'months': n, 'head_code': ps['head_code']})
        results, skipped, used = [], [], set()
        for sh in xl.sheet_names:
            ps = _parse_target_sheet(xl.parse(sh, header=None, dtype=object))
            if ps is None:
                skipped.append({'sheet': sh, 'reason': 'không có dòng Tháng'})
                continue
            sc = _store_for_sheet(cur, ps, valid, used)
            if not sc:
                skipped.append({'sheet': sh, 'reason': f"không nhận ra cửa hàng (HEAD {ps['head_code'] or '?'} - {ps['head_name'] or '?'})"})
                continue
            used.add(sc)
            n = _save_target_sheet(cur, sc, ps, actor)
            results.append({'sheet': sh, 'store': sc, 'head_code': ps['head_code'], 'head_name': ps['head_name'], 'months': n})
        if not results:
            db.rollback()
            return _json_err('Không nhận ra sheet nào là chỉ tiêu cửa hàng. ' + '; '.join(f"{x['sheet']}: {x['reason']}" for x in skipped))
        db.commit()
        cur.close()
        return jsonify({'results': results, 'skipped': skipped, 'months': sum(r['months'] for r in results)})
    except Exception as e:
        db.rollback()
        return _json_err(f'Không đọc được file: {e}')


# --- Admin: bảng tra dùng chung ------------------------------------------------------------------
def _read_ref_sheet(xl, name_hint, header_keys):
    sheet = next((s for s in xl.sheet_names if name_hint in _key(s)), None)
    if sheet is None:
        return None
    d = xl.parse(sheet, header=None, dtype=object)
    for i in range(min(12, len(d))):
        keys = [_key(x) if x is not None and not (isinstance(x, float) and np.isnan(x)) else '' for x in d.iloc[i].tolist()]
        if set(keys) & set(header_keys):
            out = d.iloc[i + 1:].copy()
            seen, cols = {}, []
            for k in keys:
                seen[k] = seen.get(k, -1) + 1
                cols.append(k if seen[k] == 0 else f'{k}__{seen[k]}')
            out.columns = cols
            return out
    return None


@chi_tieu_bp.route('/api/chi-tieu/ref/import', methods=['POST'])
def api_ref_import():
    """Tải cả file Excel chỉ tiêu (hoặc file riêng) -> nạp sheet Part Category, Hao Mòn, Oil. Thay thế bản cũ."""
    role, _ = _who()
    if role != 'admin':
        return _json_err('Chỉ admin.', 403)
    f = request.files.get('file')
    if not f:
        return _json_err('Chưa chọn file.')
    db = get_db()
    cur = db.cursor()
    try:
        xl = pd.ExcelFile(io.BytesIO(f.read()))
        done = []
        pc = _read_ref_sheet(xl, 'part category', {'ma phu tung'})
        if pc is not None:
            c_code, c_type, c_cate = _pick(pc, 'ma phu tung'), _pick(pc, 'parts type'), _pick(pc, 'phan loai')
            if c_code and c_cate:
                d = pd.DataFrame({'code': pc[c_code].map(_norm_code), 'ptype': pc[c_type] if c_type else '',
                                  'cate': pc[c_cate].map(lambda v: '' if v is None or (isinstance(v, float) and np.isnan(v)) else str(v).strip().upper())})
                d = d[d['code'] != ''].drop_duplicates('code')
                cur.execute("TRUNCATE ct_ref_part")
                execute_values(cur, "INSERT INTO ct_ref_part (code, ptype, cate) VALUES %s",
                               list(d.itertuples(index=False, name=None)), page_size=5000)
                done.append(f'Part Category: {len(d):,} mã')
        hm = _read_ref_sheet(xl, 'hao mon', {'ma phu tung'})
        if hm is not None:
            c_code, c_grp = _pick(hm, 'ma phu tung'), _pick(hm, 'nhom hao mon')
            if c_code and c_grp:
                d = pd.DataFrame({'code': hm[c_code].map(_norm_code),
                                  'grp': hm[c_grp].map(lambda v: '' if v is None or (isinstance(v, float) and np.isnan(v)) else str(v).strip())})
                d = d[(d['code'] != '') & (d['grp'] != '')].drop_duplicates('code')
                cur.execute("TRUNCATE ct_ref_hm")
                execute_values(cur, "INSERT INTO ct_ref_hm (code, grp) VALUES %s", list(d.itertuples(index=False, name=None)), page_size=5000)
                done.append(f'Hao Mòn: {len(d):,} mã')
        oil = _read_ref_sheet(xl, 'oil', {'ma hang'})
        if oil is not None:
            c_code, c_kind, c_veh = _pick(oil, 'ma hang'), _pick(oil, 'loai nhot'), _pick(oil, 'loai xe')
            if c_code:
                d = pd.DataFrame({'code': oil[c_code].map(_norm_code),
                                  'kind': oil[c_kind] if c_kind else '', 'veh': oil[c_veh] if c_veh else ''})
                d = d[d['code'] != ''].drop_duplicates('code').fillna('')
                cur.execute("TRUNCATE ct_ref_oil")
                execute_values(cur, "INSERT INTO ct_ref_oil (code, kind, vehicle) VALUES %s",
                               [(a, str(b).strip(), str(c).strip()) for a, b, c in d.itertuples(index=False, name=None)], page_size=2000)
                done.append(f'Oil: {len(d):,} mã')
        if not done:
            db.rollback()
            return _json_err('Không thấy sheet Part Category / Hao Mòn / Oil trong file.')
        _bump(cur, 'ref')
        db.commit()
        cur.close()
        return jsonify({'done': done})
    except Exception as e:
        db.rollback()
        return _json_err(f'Không đọc được file: {e}')


@chi_tieu_bp.route('/api/chi-tieu/ref/unclassified')
def api_unclassified():
    """Mã đã bán/nhận/đặt nhưng chưa có trong Part Category (Excel cũ lặng lẽ bỏ qua các mã này)."""
    role, _ = _who()
    if role != 'admin':
        return _json_err('Chỉ admin.', 403)
    month = _month_arg()
    cur = get_db().cursor()
    cur.execute("""SELECT r.part_code, SUM(r.qty*r.price) AS amount, COUNT(DISTINCT r.store_code) AS stores
                   FROM ct_rows r LEFT JOIN ct_ref_part p ON p.code = r.part_code
                   WHERE r.month=%s AND p.code IS NULL GROUP BY r.part_code ORDER BY amount DESC LIMIT 200""", (month,))
    rows = [dict(code=r['part_code'], amount=float(r['amount'] or 0), stores=r['stores']) for r in cur.fetchall()]
    cur.close()
    return jsonify({'rows': rows})


def _style_workbook(w):
    """Định dạng file Excel xuất ra: tiêu đề navy chữ trắng, số có dấu phân cách, cột tự giãn, cố định dòng đầu."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    fill = PatternFill('solid', fgColor='1E3A5F')
    side = Side(style='thin', color='D9DEE5')
    for ws in w.book.worksheets:
        ws.freeze_panes = 'B2'
        for c in ws[1]:
            c.fill = fill
            c.font = Font(bold=True, color='FFFFFF')
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = Border(bottom=side)
                if isinstance(c.value, (int, float)):
                    hdr = str(ws.cell(row=1, column=c.column).value or '')
                    c.number_format = '0.0%' if ('%' in hdr and abs(c.value) <= 1.5) else '#,##0'
        for col in ws.columns:
            m = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(m + 3, 10), 40)


# --- Xuất Excel ---------------------------------------------------------------------------------
@chi_tieu_bp.route('/api/chi-tieu/export')
def api_export():
    role, own = _who()
    month = _month_arg()
    cur = get_db().cursor()
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine='openpyxl') as w:
        if role == 'admin' and not request.args.get('store'):
            rows = [_summary_row(_result(cur, s, month)) for s in _track_stores(cur)]
            df = pd.DataFrame([{
                'Cửa hàng': r['store'], 'Chỉ tiêu Honda': r['target'], 'Kế hoạch xuất': r['plan'], 'Thực tế xuất': r['xuat'],
                '% đạt so chỉ tiêu': r['pct'], 'Đặt hàng': r['dat'], 'Nhận hàng': r['nhan'], 'Nhận/Bán': r['ratio'],
                'HM nhóm 1 - Chỉ tiêu': r['hm1_target'], 'HM nhóm 1 - Đã xuất': r['hm1_xuat'], 'HM nhóm 1 - % đạt': r['hm1_done'], 'HM nhóm 1 - Nhận/Bán': r['hm1_ratio'],
                'HM nhóm 2 - Chỉ tiêu': r['hm2_target'], 'HM nhóm 2 - Đã xuất': r['hm2_xuat'], 'HM nhóm 2 - % đạt': r['hm2_done'], 'HM nhóm 2 - Nhận/Bán': r['hm2_ratio'],
                'Cảnh báo': '; '.join(r['flags'])} for r in rows])
            df.to_excel(w, sheet_name=f'Tổng hợp {month}', index=False)
        else:
            store = _resolve_store(cur, role, own, request.args.get('store'))
            if not store:
                return _json_err('Cửa hàng không hợp lệ.')
            r = _result(cur, store, month)
            t, a, p = r['target'], r['actual'], r['plan']
            pd.DataFrame([
                ['Chỉ tiêu Honda (tổng)', t['total'], None, None],
                ['Xuất hàng', p['xuat'], a['xuat'], r['gap']['xuat']],
                ['Đặt hàng', p['dat'], a['dat'], r['gap']['dat']],
                ['Nhận hàng', p['nhan'], a['nhan'], r['gap']['nhan']],
                ['Tỉ lệ nhận/bán', None, r['ratio_nhan_ban'], None],
                ['Hao mòn nhóm 1', r['hm1']['target'], r['hm1']['xuat'], r['hm1']['target'] - r['hm1']['xuat']],
                ['Hao mòn nhóm 2', r['hm2']['target'], r['hm2']['xuat'], r['hm2']['target'] - r['hm2']['xuat']],
            ], columns=['Hạng mục', 'Mục tiêu/Kế hoạch', 'Thực tế', 'Còn thiếu']).to_excel(w, sheet_name='Kết quả', index=False)
            pd.DataFrame(r['tuan_xuat']).drop(columns='nhom').to_excel(w, sheet_name='Xuất theo tuần', index=False)
            pd.DataFrame(r['tuan_nhan']).to_excel(w, sheet_name='Nhận theo tuần', index=False)
            pd.DataFrame([{'Ngày': x['ngay'], 'Bán lẻ': x['ban_le'], 'Dịch vụ': x['dich_vu']} for x in r['ngay']]).to_excel(w, sheet_name='Theo ngày', index=False)
            try:  # chi tiết ngày theo từng nhóm (chỉ đọc dữ liệu đã tính, không đổi logic)
                rows_g = []
                for x in r['ngay']:
                    row = {'Ngày': x['ngay'], 'Bán lẻ': x['ban_le'], 'Dịch vụ': x['dich_vu']}
                    for g, v in (x.get('nhom') or {}).items():
                        row[f'{g} - BL'] = v[0]
                        row[f'{g} - DV'] = v[1]
                    rows_g.append(row)
                if rows_g:
                    pd.DataFrame(rows_g).to_excel(w, sheet_name='Ngày theo nhóm', index=False)
            except Exception:
                pass
        _style_workbook(w)
    cur.close()
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=f'chi_tieu_{month}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ============================================================================
# KPI SAU BÁN HÀNG CỦA HONDA (theo Quý / Năm)  -  file "Gửi HEAD" mỗi cửa hàng 1 file
# ----------------------------------------------------------------------------
#  * Admin nạp các file Honda gửi hàng tháng. Mỗi lần nạp THAY THẾ dữ liệu KPI cũ của cửa hàng (xoá cũ rồi ghi mới).
#  * Tháng đã có trong file Honda (<= "cập nhật đến") lấy từ file; các tháng SAU đó lấy từ dữ liệu HMS cửa hàng đang
#    tải ở mục "Tải dữ liệu" (ct_rows) rồi cộng vào quý/năm tương ứng. Honda gửi file mới -> tháng đó chuyển sang số Honda.
#  * Không sửa bất kỳ bảng/hàm cũ nào: thêm bảng ct_kpi, dùng lại _prepare/_by_cate/_by_grp/_bump/_store_for_sheet.
# ============================================================================
import json as _json
import kpi_honda

_kpi_ready = False
_kpi_lock = threading.Lock()
_kpi_cache = {}            # store -> (ver, result)


def ensure_kpi_table():
    global _kpi_ready
    if _kpi_ready:
        return
    with _kpi_lock:
        if _kpi_ready:
            return
        db = get_db()
        cur = db.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS ct_kpi (
                store_code TEXT PRIMARY KEY,
                head_code TEXT, head_name TEXT, title TEXT,
                upto CHAR(7),                      -- 'YYYY-MM' tháng Honda cập nhật đến
                data TEXT NOT NULL,                -- JSON do kpi_honda.parse_kpi_xlsx tạo
                filename TEXT, uploaded_by TEXT, uploaded_at TIMESTAMP
            );
        """)
        db.commit()
        cur.close()
        _kpi_ready = True


def _kpi_live(cur, store, months, ref):
    """Số liệu HMS của các tháng đang dùng, theo đúng định nghĩa các dòng trong file Honda."""
    out = {}
    for m in months:
        cur.execute("SELECT kind, part_code, qty, price, dt FROM ct_rows WHERE store_code=%s AND month=%s", (store, m))
        rows = cur.fetchall()
        base = pd.DataFrame(rows, columns=['kind', 'part_code', 'qty', 'price', 'dt']) if rows else \
            pd.DataFrame(columns=['kind', 'part_code', 'qty', 'price', 'dt'])
        fr = {k: _prepare(base[base['kind'] == k].drop(columns='kind'), ref) for k in KINDS}
        c_otc, c_jc, c_nhan = _by_cate(fr['otc']), _by_cate(fr['jc']), _by_cate(fr['nhan'])
        g_otc, g_jc, g_nhan = _by_grp(fr['otc']), _by_grp(fr['jc']), _by_grp(fr['nhan'])
        ban, nhan = {}, {}
        for g in HM_ALL:
            k = kpi_honda.item_key(g)
            ban[k] = g_otc[g] + g_jc[g]
            nhan[k] = g_nhan[g]
        out[m] = dict(pt=c_otc['BP'] + c_otc['GR'] + c_otc['PM'] + c_jc['BP'] + c_jc['GR'] + c_jc['PM'],
                      oil=c_otc['OIL'] + c_jc['OIL'], vap=c_otc['PG'] + c_jc['PG'],
                      recv_pt=c_nhan['BP'] + c_nhan['GR'] + c_nhan['PM'], recv_oil=c_nhan['OIL'],
                      ban=ban, nhan=nhan)
    return out


def _kpi_result(cur, store):
    cur.execute("SELECT COALESCE(MAX(ver),0) AS v FROM ct_meta WHERE key IN ('ref','kpi') OR key LIKE %s", (f's:{store}:%',))
    ver = int(cur.fetchone()['v'])
    with _cache_lock:
        hit = _kpi_cache.get(store)
    if hit and hit[0] == ver:
        return hit[1]
    cur.execute("SELECT * FROM ct_kpi WHERE store_code=%s", (store,))
    row = cur.fetchone()
    if not row:
        res = {'has': False, 'store': store}
    else:
        data = _json.loads(row['data'])
        cur.execute("SELECT DISTINCT month FROM ct_rows WHERE store_code=%s AND month > %s ORDER BY month", (store, data['upto']))
        in_file = set(data['months'])
        live_months = [r['month'] for r in cur.fetchall() if r['month'] in in_file]
        live = {}
        if live_months:
            ref_ver = _get_vers(cur, store, live_months[0])[0]
            live = _kpi_live(cur, store, live_months, _load_ref(cur, ref_ver))
        res = kpi_honda.build(data, live, NHAN_BAN_LOW, NHAN_BAN_HIGH)
        res.update(has=True, store=store, head_code=row['head_code'] or '', head_name=row['head_name'] or '',
                   file=row['filename'] or '', by=row['uploaded_by'] or '',
                   at=row['uploaded_at'].strftime('%d/%m/%Y %H:%M') if row['uploaded_at'] else '')
    res['ver'] = ver
    with _cache_lock:
        if len(_kpi_cache) > 100:
            _kpi_cache.clear()
        _kpi_cache[store] = (ver, res)
    return res


def _kpi_summary(r):
    if not r.get('has'):
        return {'store': r['store'], 'has': False}
    per = {}
    for k, p in r['periods'].items():
        per[k] = dict(range=p['range'], src=p['src'], total=p['rev']['total'], target=p['rev']['target'], rate=p['rev']['rate'],
                      recv=p['recv']['total'], ratio=p['recv']['ratio'], eval=p['recv']['eval'], prov=p['recv']['prov'],
                      g1={x: p['g1'][x] for x in ('ban', 'nhan', 'target', 'rate', 'ratio', 'eval', 'prov')},
                      g2={x: p['g2'][x] for x in ('ban', 'nhan', 'target', 'rate', 'ratio', 'eval', 'prov')})
    return {'store': r['store'], 'has': True, 'upto_label': r['upto_label'], 'next_label': r['next_label'],
            'live_labels': r['live_labels'], 'cur_q': r['cur_q'], 'periods': per, 'file': r['file'], 'at': r['at']}


@chi_tieu_bp.route('/api/chi-tieu/kpi/data')
def api_kpi_data():
    role, own = _who()
    ensure_kpi_table()
    cur = get_db().cursor()
    store = _resolve_store(cur, role, own, request.args.get('store'))
    if not store:
        return _json_err('Vui lòng chọn cửa hàng.')
    res = _kpi_result(cur, store)
    cur.close()
    return jsonify(res)


@chi_tieu_bp.route('/api/chi-tieu/kpi/overview')
def api_kpi_overview():
    role, _ = _who()
    if role != 'admin':
        return _json_err('Chỉ admin xem được tổng hợp.', 403)
    ensure_kpi_table()
    cur = get_db().cursor()
    rows = [_kpi_summary(_kpi_result(cur, s)) for s in _track_stores(cur)]
    cur.close()
    have = [r for r in rows if r['has']]
    return jsonify({'rows': rows, 'cur_q': have[0]['cur_q'] if have else None,
                    'upto_label': have[0]['upto_label'] if have else None,
                    'next_label': have[0]['next_label'] if have else None,
                    'ranges': {k: v['range'] for k, v in have[0]['periods'].items()} if have else {}})


@chi_tieu_bp.route('/api/chi-tieu/kpi/version')
def api_kpi_version():
    role, own = _who()
    ensure_kpi_table()
    cur = get_db().cursor()
    if role == 'admin' and not request.args.get('store'):
        cur.execute("SELECT COALESCE(MAX(ver),0) AS v FROM ct_meta WHERE key IN ('ref','kpi') OR key LIKE 's:%'")
    else:
        store = _resolve_store(cur, role, own, request.args.get('store'))
        if not store:
            return _json_err('Cửa hàng không hợp lệ.')
        cur.execute("SELECT COALESCE(MAX(ver),0) AS v FROM ct_meta WHERE key IN ('ref','kpi') OR key LIKE %s", (f's:{store}:%',))
    v = cur.fetchone()['v']
    cur.close()
    return jsonify({'v': int(v)})


@chi_tieu_bp.route('/api/chi-tieu/kpi/upload', methods=['POST'])
def api_kpi_upload():
    """Admin nạp file KPI Honda (nhiều file một lúc, mỗi file 1 cửa hàng; tự nhận cửa hàng theo mã HEAD / tên '#N').
    Dữ liệu KPI cũ của cửa hàng được xoá rồi ghi mới. replace_all=1: cửa hàng không có trong lần nạp này cũng bị xoá KPI cũ."""
    role, _ = _who()
    if role != 'admin':
        return _json_err('Chỉ admin.', 403)
    ensure_kpi_table()
    files = [f for f in request.files.getlist('files') if f and f.filename]
    if not files:
        return _json_err('Chưa chọn file.')
    replace_all = request.form.get('replace_all') == '1'
    db = get_db()
    cur = db.cursor()
    valid = _valid_store_codes(cur)
    actor = session.get('full_name') or session.get('user')
    parsed, errors, used = [], [], set()
    for f in files:
        try:
            data = kpi_honda.parse_kpi_xlsx(f.read())
        except Exception as e:
            errors.append(f'{f.filename}: {e}')
            continue
        sc = _store_for_sheet(cur, {'head_code': data['head_code'], 'head_name': data['head_name']}, valid, used)
        if not sc:
            errors.append(f"{f.filename}: không nhận ra cửa hàng (HEAD {data['head_code'] or '?'} - {data['head_name'] or '?'})")
            continue
        used.add(sc)
        parsed.append((sc, f.filename, data))
    if not parsed:
        cur.close()
        return _json_err('Không nạp được file nào: ' + '; '.join(errors))
    try:
        if replace_all:
            cur.execute("DELETE FROM ct_kpi WHERE NOT (store_code = ANY(%s))", (sorted(used),))
        done = []
        for sc, fname, data in parsed:
            cur.execute("DELETE FROM ct_kpi WHERE store_code=%s", (sc,))
            cur.execute("""INSERT INTO ct_kpi (store_code, head_code, head_name, title, upto, data, filename, uploaded_by, uploaded_at)
                           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                        (sc, data['head_code'], data['head_name'], data['title'], data['upto'],
                         _json.dumps(data, ensure_ascii=False), fname, actor, vn_now()))
            if data['head_code']:
                cur.execute("""INSERT INTO ct_stores (store_code, head_code, head_name) VALUES (%s,%s,%s)
                               ON CONFLICT (store_code) DO UPDATE SET head_code=COALESCE(ct_stores.head_code, EXCLUDED.head_code),
                               head_name=COALESCE(ct_stores.head_name, EXCLUDED.head_name)""",
                            (sc, data['head_code'], data['head_name']))
            done.append(f"{sc} ← {fname} (đến {kpi_honda._ml(data['upto'])})")
        _bump(cur, 'kpi')
        db.commit()
        with _cache_lock:
            _kpi_cache.clear()
    except Exception as e:
        db.rollback()
        cur.close()
        return _json_err(f'Lỗi lưu dữ liệu KPI ({e})', 500)
    cur.close()
    return jsonify({'done': done, 'errors': errors})