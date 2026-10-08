# -*- coding: utf-8 -*-
"""
KPI SAU BÁN HÀNG của Honda (file "Gửi HEAD", mỗi cửa hàng 1 file) -> đọc file + gộp với dữ liệu HMS của tháng đang dùng.

Module thuần Python (không phụ thuộc Flask/DB) để dễ kiểm thử. Route nằm trong chi_tieu.py.

Cấu trúc file Honda (đã kiểm tra trên 5 file mẫu):
  * B1 = Code HEAD (53001...), B2 = Tên HEAD ("Nam Sương #1"), C1 chứa "CẬP NHẬT ĐẾN T9/2026".
  * Dòng tiêu đề: cột A = "Kết quả kinh doanh", các cột T3..T3 (13 tháng, dòng phía trên là năm),
    rồi "Quý 1..Quý 4" và "Cả Ki".
  * Quý của Honda KHÔNG theo lịch: Q1 = T3-T6, Q2 = T7-T9, Q3 = T10-T12, Q4 = T1-T2 (T3 năm sau thuộc kỳ mới).
    Module tự suy ra tháng của từng quý bằng cách cộng dòng "Mục tiêu" cho khớp mục tiêu quý trong file.
  * Các số trong file là số cố định (không phải công thức) nên đọc thẳng được.
"""
import io
import re
import unicodedata

import openpyxl

NHAN_BAN_LOW, NHAN_BAN_HIGH = 0.95, 1.05
DEFAULT_QMONTHS = {'1': (3, 4, 5, 6), '2': (7, 8, 9), '3': (10, 11, 12), '4': (1, 2)}   # theo tháng lịch


def _strip_accents(s):
    s = unicodedata.normalize('NFD', str(s))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return s.replace('đ', 'd').replace('Đ', 'D')


def key(s):
    return re.sub(r'[^a-z0-9]+', ' ', _strip_accents(s).lower()).strip()


def item_key(name):
    """Khoá tên hạng mục hao mòn: bỏ dấu, bỏ khoảng trắng/gạch ngang ('Bu-gi' == 'Bugi')."""
    return key(name).replace(' ', '')


def _blank(v):
    return v is None or (isinstance(v, float) and v != v) or (isinstance(v, str) and not v.strip())


def _num(v):
    if _blank(v) or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '').strip())
    except ValueError:
        return None


def _ym(y, m):
    return f'{int(y):04d}-{int(m):02d}'


def _next_month(ym):
    y, m = int(ym[:4]), int(ym[5:7])
    return _ym(y + (m == 12), 1 if m == 12 else m + 1)


# ----------------------------------------------------------------------------
# Đọc file
# ----------------------------------------------------------------------------
def parse_kpi_xlsx(raw):
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = None
    for s in wb.worksheets:                       # ưu tiên sheet có "Code HEAD" ở cột A
        for r in s.iter_rows(min_row=1, max_row=10, max_col=1, values_only=True):
            if r[0] is not None and key(r[0]) == 'code head':
                ws = s
                break
        if ws:
            break
    ws = ws or wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    width = max((len(r) for r in rows), default=0)
    rows = [r + [None] * (width - len(r)) for r in rows]

    head_code = head_name = None
    title = ''
    for r in rows[:10]:
        k0 = key(r[0]) if not _blank(r[0]) else ''
        if k0 == 'code head' and not _blank(r[1]):
            v = r[1]
            head_code = str(int(v)) if isinstance(v, float) and v.is_integer() else re.sub(r'\s+', '', str(v)).upper()
        elif k0 == 'ten head' and not _blank(r[1]):
            head_name = str(r[1]).strip()
    if rows and not _blank(rows[0][2]):
        title = re.sub(r'\s+', ' ', str(rows[0][2])).strip()

    # --- dòng tiêu đề ---
    h0 = next((i for i, r in enumerate(rows) if not _blank(r[0]) and key(r[0]) == 'ket qua kinh doanh'), None)
    if h0 is None or h0 == 0:
        raise ValueError('Không đúng định dạng file KPI Honda (không thấy dòng "Kết quả kinh doanh").')
    months, qcol, kcol = {}, {}, None
    for ci, v in enumerate(rows[h0]):
        if _blank(v):
            continue
        s = str(v).strip()
        m = re.fullmatch(r'T\s*(\d{1,2})', s)
        if m:
            y = rows[h0 - 1][ci]
            if _blank(y):
                continue
            months[ci] = _ym(int(float(y)), int(m.group(1)))
            continue
        kq = key(s)
        m = re.fullmatch(r'quy (\d)', kq)
        if m:
            qcol[m.group(1)] = ci
        elif kq.startswith('ca k'):
            kcol = ci
    if not months:
        raise ValueError('Không đọc được các cột tháng (T3, T4...) của file KPI Honda.')
    month_list = [months[c] for c in sorted(months)]

    # --- các dòng số liệu ---
    series, items, order = {}, {'1': [], '2': []}, []

    def put(sid, r):
        s = {'m': {}, 'q': {}, 'k': None}
        for ci, ym in months.items():
            v = r[ci]
            s['m'][ym] = (str(v).strip().upper() if isinstance(v, str) else _num(v))
        for q, ci in qcol.items():
            v = r[ci]
            s['q'][q] = (str(v).strip().upper() if isinstance(v, str) else _num(v))
        if kcol is not None:
            v = r[kcol]
            s['k'] = (str(v).strip().upper() if isinstance(v, str) else _num(v))
        series[sid] = s

    cur_a, grp = '', None
    for i in range(h0 + 1, len(rows)):
        r = rows[i]
        if not _blank(r[0]):
            ka = key(r[0])
            if ka == 'ket qua kinh doanh':
                continue
            cur_a = ka
            m = re.search(r'nhom (\d)', cur_a)
            if m:
                grp = m.group(1)
            elif cur_a.startswith('ti le nhan ban') or cur_a.startswith('doanh thu') or cur_a.startswith('tong doanh thu'):
                grp = None
        if _blank(r[2]):
            continue
        raw_c = str(r[2]).strip()
        c = key(raw_c)
        if grp is None:                                                    # khối tổng
            if cur_a.startswith('doanh thu pt') and c == 'thuc te':
                put('pt', r)
            elif cur_a.startswith('doanh thu dau') and c == 'thuc te':
                put('oil', r)
            elif cur_a.startswith('doanh thu vap') and c == 'thuc te':
                put('vap', r)
            elif cur_a.startswith('tong doanh thu'):
                sid = {'thuc te': 'total', 'muc tieu': 'target', 'ti le dat': 'rate'}.get(c)
                if sid:
                    put(sid, r)
            elif cur_a.startswith('ti le nhan ban'):
                sid = {'phu tung': 'recv_pt', 'dau': 'recv_oil', 'tong nhan': 'recv_total',
                       'nhan ban': 'recv_ratio', 'danh gia': 'eval'}.get(c)
                if sid:
                    put(sid, r)
            continue
        g = f'g{grp}'
        ik = item_key(raw_c)
        if 'doanh thu thuc te' in cur_a:                                   # từng hạng mục: bán
            if not any(x['key'] == ik for x in items[grp]):
                items[grp].append({'key': ik, 'name': raw_c})
            put(f'{g}.i.{ik}.ban', r)
        elif re.fullmatch(r'nhom \d nhan', cur_a):                         # từng hạng mục: nhận
            if not any(x['key'] == ik for x in items[grp]):
                items[grp].append({'key': ik, 'name': raw_c})
            put(f'{g}.i.{ik}.nhan', r)
        elif cur_a.startswith('tong nhan nhom'):
            if c == 'nhan':
                put(f'{g}.nhan', r)
        elif cur_a.startswith('tong ban nhom'):
            continue
        elif cur_a.startswith('ti le nhan ban nhom'):
            if c == 'nhan ban':
                put(f'{g}.ratio', r)
        elif cur_a == 'danh gia':
            if c == 'ok ng':
                put(f'eval_{g}', r)
        elif cur_a.startswith('nhom'):                                     # dòng tổng nhóm: DT thực tế / mục tiêu / tỉ lệ đạt
            sid = {'dt thuc te': f'{g}.ban', 'dt muc tieu': f'{g}.target', 'ti le dat': f'{g}.rate'}.get(c)
            if sid:
                put(sid, r)

    for need in ('total', 'target'):
        if need not in series:
            raise ValueError('Không đúng định dạng file KPI Honda (thiếu dòng tổng doanh thu / mục tiêu).')
    if all((v or 0) == 0 for v in series['total']['m'].values()) and all((v or 0) == 0 for v in series['target']['m'].values()):
        raise ValueError('File chưa có số liệu (có thể file lưu công thức chưa tính). Mở bằng Excel, lưu lại rồi nạp lại.')

    # --- tháng cập nhật đến ---
    upto = None
    txt = ' '.join(str(rows[i][2]) for i in range(min(4, len(rows))) if not _blank(rows[i][2]))
    m = re.search(r'T\s*(\d{1,2})\s*/\s*(\d{4})', txt)
    if m and _ym(m.group(2), m.group(1)) in month_list:
        upto = _ym(m.group(2), m.group(1))
    if upto is None:                                                       # dự phòng: tháng cuối còn có doanh thu
        done = [ym for ym in month_list if (series['total']['m'].get(ym) or 0) > 0]
        upto = done[-1] if done else month_list[0]

    # --- tháng của từng quý (suy ra từ mục tiêu) ---
    qmonths = _derive_qmonths(month_list, series['target'])

    return {
        'title': title, 'head_code': head_code, 'head_name': head_name, 'upto': upto,
        'months': month_list, 'qmonths': qmonths, 'series': series,
        'items': {k: v for k, v in items.items() if v},
    }


def _derive_qmonths(month_list, tgt):
    t = tgt['m']
    out, pos, ok = {}, 0, True
    for q in ('1', '2', '3', '4'):
        qt = tgt['q'].get(q)
        if qt is None:
            ok = False
            break
        acc, run, found = 0.0, [], False
        for ym in month_list[pos:]:
            acc += t.get(ym) or 0.0
            run.append(ym)
            if abs(acc - qt) <= max(1.0, abs(qt) * 1e-6):
                found = True
                break
        if not found:
            ok = False
            break
        out[q] = run
        pos += len(run)
    if ok:
        return out
    out = {}                                                               # dự phòng theo tháng lịch
    for q, ms in DEFAULT_QMONTHS.items():
        out[q] = [ym for ym in month_list if int(ym[5:7]) in ms and ym < _ym(int(month_list[0][:4]) + 1, int(month_list[0][5:7]))]
    return out


# ----------------------------------------------------------------------------
# Gộp Honda (tháng đã chốt) + HMS (tháng đang dùng) -> quý / năm
# ----------------------------------------------------------------------------
def _ml(ym):
    return f'T{int(ym[5:7])}/{ym[:4]}'


def _range(ms):
    if not ms:
        return ''
    a, b = ms[0], ms[-1]
    if a == b:
        return _ml(a)
    if a[:4] == b[:4]:
        return f'T{int(a[5:7])}–T{int(b[5:7])}/{a[:4]}'
    return f'{_ml(a)}–{_ml(b)}'


def build(data, live=None, lo=NHAN_BAN_LOW, hi=NHAN_BAN_HIGH):
    """data: kết quả parse_kpi_xlsx (đọc lại từ DB).
    live: {'YYYY-MM': {pt, oil, vap, recv_pt, recv_oil, ban: {item_key: số}, nhan: {item_key: số}}} cho các tháng SAU 'upto'."""
    live = live or {}
    S, upto, months = data['series'], data['upto'], data['months']
    qmonths = data['qmonths']
    items = data.get('items') or {}

    def sv(sid, ym):
        s = S.get(sid)
        v = s['m'].get(ym) if s else None
        return v if isinstance(v, (int, float)) else 0.0

    def src_of(ym):
        if ym <= upto:
            return 'honda'
        return 'hms' if ym in live else 'none'

    def month_vals(ym):
        src = src_of(ym)
        o = dict(src=src, pt=0.0, oil=0.0, vap=0.0, total=0.0, recv_pt=0.0, recv_oil=0.0, recv_total=0.0, g={})
        for gk in items:
            o['g'][gk] = {x['key']: [0.0, 0.0] for x in items[gk]}
        if src == 'honda':
            o.update(pt=sv('pt', ym), oil=sv('oil', ym), vap=sv('vap', ym), total=sv('total', ym),
                     recv_pt=sv('recv_pt', ym), recv_oil=sv('recv_oil', ym), recv_total=sv('recv_total', ym))
            for gk in items:
                for x in items[gk]:
                    o['g'][gk][x['key']] = [sv(f'g{gk}.i.{x["key"]}.ban', ym), sv(f'g{gk}.i.{x["key"]}.nhan', ym)]
        elif src == 'hms':
            L = live[ym]
            o.update(pt=L['pt'], oil=L['oil'], vap=L['vap'], total=L['pt'] + L['oil'] + L['vap'],
                     recv_pt=L['recv_pt'], recv_oil=L['recv_oil'], recv_total=L['recv_pt'] + L['recv_oil'])
            for gk in items:
                for x in items[gk]:
                    o['g'][gk][x['key']] = [L['ban'].get(x['key'], 0.0), L['nhan'].get(x['key'], 0.0)]
        return o

    mv = {ym: month_vals(ym) for ym in months}

    def honda_eval(sid, scope, ident):
        s = S.get(sid)
        if not s:
            return None
        v = s['m'].get(ident) if scope == 'm' else (s['q'].get(ident) if scope == 'q' else s['k'])
        return v if v in ('OK', 'NG') else None

    def judge(ratio, has_data, live_contrib, sid, scope, ident):
        """Đánh giá theo chuẩn nhận/bán lo..hi (95-105%): trong khoảng = OK, ngoài khoảng = NG (không dùng chữ của Honda)."""
        if not has_data or ratio is None:
            return None, False
        return ('OK' if lo <= ratio <= hi else 'NG'), False

    def period(ms, scope, ident):
        vals = [mv[m] for m in ms]
        used = [v for v in vals if v['src'] != 'none']
        has = bool(used)
        live_c = any(v['src'] == 'hms' for v in vals)
        tot = {k: sum(v[k] for v in used) for k in ('pt', 'oil', 'vap', 'total', 'recv_pt', 'recv_oil', 'recv_total')}
        target = sum(sv('target', m) for m in ms)
        rate = (tot['total'] / target) if (has and target) else None
        ratio = (tot['recv_total'] / tot['total']) if (has and tot['total']) else None
        ev, prov = judge(ratio, has, live_c, 'eval', scope, ident)
        if not has:
            src = 'none'
        elif live_c and any(v['src'] == 'honda' for v in vals):
            src = 'mixed'
        elif live_c:
            src = 'hms'
        else:
            src = 'honda'
        out = dict(months=ms, range=_range(ms), src=src, open=any(v['src'] == 'none' for v in vals),
                   rev=dict(pt=tot['pt'], oil=tot['oil'], vap=tot['vap'], total=tot['total'], target=target, rate=rate),
                   recv=dict(pt=tot['recv_pt'], oil=tot['recv_oil'], total=tot['recv_total'], ratio=ratio, eval=ev, prov=prov))
        for gk in items:
            its, ban, nhan = [], 0.0, 0.0
            for x in items[gk]:
                b = sum(v['g'][gk][x['key']][0] for v in used)
                n = sum(v['g'][gk][x['key']][1] for v in used)
                its.append(dict(key=x['key'], name=x['name'], ban=b, nhan=n))
                ban += b
                nhan += n
            gt = sum(sv(f'g{gk}.target', m) for m in ms)
            gr = (nhan / ban) if (has and ban) else None
            gev, gprov = judge(gr, has, live_c, f'eval_g{gk}', scope, ident)
            out[f'g{gk}'] = dict(ban=ban, nhan=nhan, target=gt, rate=(ban / gt) if (has and gt) else None,
                                 ratio=gr, eval=gev, prov=gprov, items=its)
        return out

    periods = {q: period(qmonths[q], 'q', q) for q in ('1', '2', '3', '4') if q in qmonths}
    all_ms = [m for q in ('1', '2', '3', '4') for m in qmonths.get(q, [])]
    periods['all'] = period(all_ms, 'k', None)
    for q in ('1', '2', '3', '4'):
        if q in periods:
            periods[q]['label'] = f'Quý {q}'
    periods['all']['label'] = 'Cả năm'

    monthly = []
    for ym in months:
        p = period([ym], 'm', ym)
        p['month'] = ym
        p['label'] = _ml(ym)
        p['in_year'] = ym in all_ms
        monthly.append(p)

    nxt = _next_month(upto)
    q_of = {m: q for q, ms in qmonths.items() for m in ms}
    cur_q = q_of.get(nxt) or q_of.get(upto) or 'all'
    return dict(title=data.get('title') or '', upto=upto, upto_label=_ml(upto), next_month=nxt,
                next_label=_ml(nxt), live_months=sorted(live), live_labels=[_ml(m) for m in sorted(live)],
                cur_q=cur_q, periods=periods, monthly=monthly,
                groups={gk: dict(names=', '.join(x['name'] for x in items[gk])) for gk in items})
