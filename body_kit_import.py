# -*- coding: utf-8 -*-
"""
Đọc sheet "Bô nhựa các đời xe " (CHÚ Ý: có dấu cách ở cuối tên sheet trong
file gốc) trong file BẢNG GIÁ BỘ ÁO XE HONDA và chuẩn hoá thành danh sách
"nhóm" (mỗi nhóm = 1 dòng xe + đời + màu cụ thể, kèm ảnh + danh sách phụ
tùng bộ áo của đúng màu/đời đó).

CẤU TRÚC GỐC (đã kiểm tra thực tế trên file mẫu):
    Cột A "Mã loại"   - nhãn nhóm, vd "K57 Balde số 2019 màu Trắng - NHB55"
                        được MERGE DỌC xuyên suốt toàn bộ khối dòng của
                        nhóm đó (mỗi khối trung bình ~15-30 dòng).
    Cột B "Model"     - mã model ngắn (vd "B2") - cũng merge dọc y hệt cột A.
    Cột C "STT"       - số thứ tự phụ tùng TRONG nhóm (reset lại từ 1 ở
                        mỗi nhóm mới).
    Cột D "Mã chi tiết" - mã phụ tùng.
    Cột E "Thay thế"  - ghi chú thay thế (vd "Dùng chung").
    Cột F "Tên tiếng việt" - tên phụ tùng.
    Cột G "Gía Honda" - giá phụ tùng chính hãng Honda.
    Cột H "Giá CRM"   - giá phụ tùng theo hệ CRM (có thể khác giá Honda).
    Cột I "TỒN KHO"   - tồn kho phụ tùng đó tại thời điểm lập file gốc.
    Cột J "Ghi chú"   - ghi chú thêm.
    Cột K:N "Hình ảnh đời xe" - MERGE thành 1 khối ảnh lớn (1 ảnh xe cho cả
                        nhóm). Ảnh được đọc TRỰC TIẾP từ cấu trúc XML/zip
                        của file .xlsx (xem _build_image_row_map bên dưới),
                        KHÔNG dùng `ws._images` của openpyxl.
    NGAY SAU dòng phụ tùng cuối cùng của 1 nhóm luôn có 1 DÒNG TỔNG (STT và
    Mã chi tiết đều trống, cột G/H là TỔNG giá cả bộ, cột I thường là chuỗi
    "0" vô nghĩa/lỗi công thức #N/A ở cột F) - dòng này KHÔNG phải phụ tùng,
    được nhận diện và gộp thành total_honda_price/total_crm_price của nhóm,
    không tạo dòng phụ tùng cho nó.

ĐỘ CHÍNH XÁC:
- Dùng merge-map cho CỘT A và CỘT B (giống bo_import.py) để "trải" giá trị
  merge dọc xuống mọi dòng thuộc khối, xác định đúng ranh giới nhóm (dòng
  nào group_label đổi khác dòng trước = nhóm mới bắt đầu).
- Ảnh được ghép vào đúng nhóm bằng cách so khớp DÒNG NEO (anchor row) của
  ảnh với dòng bắt đầu (min_row) của khối merge cột A/K tương ứng.

LƯU Ý QUAN TRỌNG VỀ CÁCH ĐỌC ẢNH (đã đổi cách làm - xem lịch sử bug):
    Bản đầu tiên của module này dùng `ws._images` (thuộc tính nội bộ của
    openpyxl) để lấy danh sách ảnh nhúng trong sheet. Cách này chạy ĐÚNG
    trên một số máy (vd. openpyxl 3.1.x) nhưng lại trả về DANH SÁCH RỖNG
    trên môi trường production thực tế (đã xác nhận qua log: 0/749 nhóm có
    ảnh dù cùng 1 file Excel, trong khi chạy local ra 743/749) - nhiều khả
    năng do file gốc được xuất từ WPS Office (có namespace "dbsheet" lạ
    trong workbook.xml, không phải Excel gốc của Microsoft) khiến các bản
    openpyxl khác nhau xử lý cấu trúc drawing/anchor không nhất quán.

    Để không còn phụ thuộc hành vi (có thể đổi giữa các bản) của
    `ws._images`, module này giờ đọc ảnh TRỰC TIẾP từ cấu trúc XML bên
    trong file .xlsx (vốn chỉ là 1 file zip theo chuẩn OOXML mở, không phụ
    thuộc thư viện đọc Excel nào):
        xl/workbook.xml                      -> tìm r:id của sheet theo tên
        xl/_rels/workbook.xml.rels           -> r:id -> đường dẫn sheetN.xml
        xl/worksheets/_rels/sheetN.xml.rels  -> tìm rId của phần "drawing"
        xl/drawings/drawingM.xml             -> mỗi ảnh là 1 anchor, có
                                                 <xdr:from><xdr:row> (dòng
                                                 neo, 0-based) và rId ảnh
        xl/drawings/_rels/drawingM.xml.rels  -> rId -> đường dẫn file ảnh
                                                 thật trong xl/media/...
    Cách này dùng CHỈ thư viện chuẩn (zipfile + xml.etree), không dùng bất
    kỳ API nội bộ nào của openpyxl nên cho kết quả giống hệt nhau trên mọi
    máy/mọi phiên bản thư viện.
"""
import posixpath
import re
import unicodedata
import zipfile
from xml.etree import ElementTree as ET

import openpyxl

SHEET_NAME = 'Bô nhựa các đời xe '  # LƯU Ý dấu cách ở cuối - có thật trong file gốc

# ----------------------------------------------------------------------------
# PHÂN LOẠI NHÓM THEO "DÒNG XE" (vd Future, Wave, Air Blade...) + "ĐỜI/KIỂU"
# (vd Future 125, Future I, Future neo...) - để hiển thị menu 2 tầng thay vì
# 1 danh sách phẳng ~750 nhóm. Suy luận TỰ ĐỘNG từ text của "Mã loại" (cột A)
# bằng khớp từ khoá dòng xe đã biết (KHÔNG có cột riêng cho việc này trong
# file Excel gốc, mọi thứ đều gộp chung vào 1 câu mô tả tự do) - vì vậy đây
# LUÔN LÀ SUY LUẬN TỐT NHẤT CÓ THỂ (best-effort), không phải trích xuất 100%
# chính xác từ dữ liệu có cấu trúc. Nhóm nào không khớp từ khoá nào sẽ rơi
# vào dòng xe "Khác" - người dùng có thể tự soát lại tại đó.
#
# Thứ tự trong danh sách QUAN TRỌNG: khớp từ trên xuống, ưu tiên độ dài khớp
# xa nhất về đầu chuỗi (xem hàm classify_group_label) - vd "ari blade" phải
# đứng trước để không bị nuốt nhầm bởi 1 từ khoá ngắn hơn ở vị trí sau.
_FAMILY_PATTERNS = [
    (r'\bari\s*blade\b', 'Air Blade'),
    (r'\bair\s*blade\b', 'Air Blade'),
    # "Balde số ..." - "số" = XE SỐ (xe côn/tay côn, vd K57...) - đây thực ra
    # là dòng "Blade" (xe số), KHÔNG PHẢI Air Blade (xe tay ga) dù cũng gõ
    # tắt/gõ nhầm thành "Balde" giống nhau trong file gốc. Phải đứng TRƯỚC
    # pattern "\bbalde\b" chung bên dưới (2 pattern khớp cùng vị trí bắt đầu
    # "balde" -> thuật toán chọn theo vị trí sớm nhất, bằng nhau thì giữ
    # pattern gặp TRƯỚC trong danh sách này, nên phải xếp trước).
    (r'\bbalde\s+so\b', 'Blade'),
    (r'\bbalde\b', 'Air Blade'),          # "Balde" (không kèm "số") - viết tắt/gõ nhầm phổ biến của Air Blade
    (r'\bblade\b', 'Blade'),              # Honda Blade - dòng xe RIÊNG, KHÁC Air Blade (đứng SAU 2 pattern
                                           # "air/ari blade" ở trên để không nuốt nhầm - "air blade"/"ari blade"
                                           # luôn khớp sớm hơn trong chuỗi nên vẫn thắng, "blade" đứng 1 mình
                                           # (không có "air"/"ari" phía trước) mới rơi vào đây).
    (r'\bsuper\s*dream\b', 'Dream'),
    (r'\bdream\b', 'Dream'),
    (r'\bfuture\b', 'Future'),
    (r'wave\b', 'Wave'),                  # không ép \b đầu để vẫn khớp cả trường hợp dính liền vd "K03WAVE"
    (r'\blead\b', 'Lead'),
    (r'\bvis(?:i|o)+nn?\b', 'Vision'),    # bắt cả biến thể gõ nhầm "Visonn"
    (r'\bpcx\b', 'PCX'),
    (r'\bclick\b', 'Click'),
    (r'\bwinner\b', 'Winner'),
    # "\bsh\b" một mình không khớp được kiểu viết liền không cách như
    # "SH350" (không có ranh giới từ giữa "h" và "3" vì cả 2 đều là ký tự
    # "chữ/số" theo \b) - thêm nhánh \bsh(?=\d) để vẫn bắt được trường hợp
    # này (chỉ cần match bắt đầu tại "sh", không cần "nuốt" luôn số phía sau).
    (r'\bsh\b|\bsh(?=\d)', 'SH'),
]
# Cắt phần "đời/kiểu" tại điểm sớm nhất xuất hiện từ khoá năm ("năm"/"đời")
# HOẶC 1 số 4 chữ số dạng năm (19xx/20xx) - vì 1 số nhãn không có từ "năm"
# mà năm đứng liền ngay sau tên đời xe (vd "K57 Balde số 2014...").
_CUTOFF_RE = re.compile(r'\b(nam|doi|mau)\b|(19|20)\d{2}', re.IGNORECASE)

# Số năm SẢN XUẤT (19xx/20xx) đầu tiên xuất hiện trong nhãn - dùng để SẮP
# XẾP các nhóm trong cùng 1 dòng xe theo đời từ CŨ -> MỚI (thấp nhất lên
# trên, cao nhất xuống dưới). Nhãn nào không tìm thấy năm nào sẽ có year =
# None, xếp xuống cuối (xem ORDER BY ở body_kit.py).
_YEAR_RE = re.compile(r'\b(19|20)\d{2}\b')


def _extract_year(label):
    if not label:
        return None
    m = _YEAR_RE.search(_strip_accents(label))
    return int(m.group(0)) if m else None


def _strip_accents(s):
    # NFD chỉ bóc được các dấu thanh/nguyên âm (kết hợp base + combining
    # mark) - chữ "đ/Đ" tiếng Việt là 1 KÝ TỰ RIÊNG (không phải "d" + dấu
    # gạch ngang dạng combining mark) nên KHÔNG bị NFD bóc, phải thay tay.
    s = s.replace('đ', 'd').replace('Đ', 'D')
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def classify_group_label(label):
    """Suy luận (vehicle_family, sub_model) từ nhãn nhóm gốc (cột "Mã
    loại"). Trả về ('Khác', label) nếu không khớp từ khoá dòng xe nào đã
    biết - xem giải thích chi tiết ở khối comment phía trên."""
    if not label:
        return ('Khác', label or '')
    plain = _strip_accents(label).lower()
    best = None  # (start_index, family)
    for pattern, family in _FAMILY_PATTERNS:
        m = re.search(pattern, plain)
        if m and (best is None or m.start() < best[0]):
            best = (m.start(), family)
    if not best:
        return ('Khác', re.sub(r'\s+', ' ', label).strip())
    start_idx, family = best
    tail = label[start_idx:]
    cutoff = _CUTOFF_RE.search(_strip_accents(tail).lower())
    sub_model = tail[:cutoff.start()] if cutoff else tail
    sub_model = re.sub(r'\s+', ' ', sub_model).strip(' -–')
    if not sub_model:
        sub_model = family
    return (family, sub_model)

# Các định dạng ảnh mà trình duyệt hiển thị trực tiếp được - .wmf/.emf (ảnh
# vẽ vector cũ của Windows) KHÔNG hiển thị được trên web nên bị loại, coi
# như nhóm đó không có ảnh thay vì lưu 1 file ảnh không xem được.
_WEB_SAFE_IMAGE_FORMATS = {'jpeg', 'jpg', 'png', 'gif', 'bmp'}

_NS = {
    'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
}
_R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _find_sheet_xml_path(zf, sheet_name):
    """Trả về đường dẫn (vd 'xl/worksheets/sheet2.xml') của sheet có tên
    chính xác `sheet_name`, hoặc None nếu không tìm thấy."""
    wb_root = ET.fromstring(zf.read('xl/workbook.xml'))
    sheets_el = wb_root.find('main:sheets', _NS)
    if sheets_el is None:
        return None
    rid = None
    for sheet_el in sheets_el:
        if sheet_el.get('name') == sheet_name:
            rid = sheet_el.get(f'{{{_R_NS}}}id')
            break
    if not rid:
        return None
    rels_root = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
    for rel in rels_root:
        if rel.get('Id') == rid:
            return 'xl/' + rel.get('Target').lstrip('/')
    return None


def _find_drawing_xml_path(zf, sheet_xml_path):
    """Trả về đường dẫn drawingM.xml được tham chiếu bởi sheet, hoặc None
    nếu sheet đó không có ảnh/drawing nào."""
    sheet_dir, sheet_file = sheet_xml_path.rsplit('/', 1)
    rels_path = f'{sheet_dir}/_rels/{sheet_file}.rels'
    if rels_path not in zf.namelist():
        return None
    rels_root = ET.fromstring(zf.read(rels_path))
    for rel in rels_root:
        if rel.get('Type', '').endswith('/drawing'):
            return posixpath.normpath(posixpath.join(sheet_dir, rel.get('Target')))
    return None


def _build_image_row_map_from_zip(path, sheet_name):
    """Đọc ảnh trực tiếp từ cấu trúc XML/zip của file .xlsx (xem giải thích
    chi tiết ở docstring đầu file) - KHÔNG dùng ws._images của openpyxl.
    Trả về dict {anchor_row (1-indexed, khớp hệ đánh số dòng của
    openpyxl): (image_bytes, format_ext)}."""
    image_map = {}
    with zipfile.ZipFile(path) as zf:
        sheet_xml_path = _find_sheet_xml_path(zf, sheet_name)
        if not sheet_xml_path:
            return image_map
        drawing_xml_path = _find_drawing_xml_path(zf, sheet_xml_path)
        if not drawing_xml_path:
            return image_map

        drawing_dir, drawing_file = drawing_xml_path.rsplit('/', 1)
        drawing_rels_path = f'{drawing_dir}/_rels/{drawing_file}.rels'
        rid_to_target = {}
        if drawing_rels_path in zf.namelist():
            rels_root = ET.fromstring(zf.read(drawing_rels_path))
            for rel in rels_root:
                rid_to_target[rel.get('Id')] = rel.get('Target')

        drawing_root = ET.fromstring(zf.read(drawing_xml_path))
        anchors = list(drawing_root.findall('xdr:twoCellAnchor', _NS)) + \
            list(drawing_root.findall('xdr:oneCellAnchor', _NS))

        for anchor in anchors:
            from_el = anchor.find('xdr:from', _NS)
            row_el = from_el.find('xdr:row', _NS) if from_el is not None else None
            if row_el is None or row_el.text is None:
                continue
            blip = anchor.find(f'.//xdr:blipFill/a:blip', _NS)
            if blip is None:
                continue
            embed_rid = blip.get(f'{{{_R_NS}}}embed')
            if not embed_rid or embed_rid not in rid_to_target:
                continue
            media_path = posixpath.normpath(posixpath.join(drawing_dir, rid_to_target[embed_rid]))
            ext = media_path.rsplit('.', 1)[-1].lower() if '.' in media_path else ''
            if ext not in _WEB_SAFE_IMAGE_FORMATS:
                continue  # .wmf/.emf hoặc định dạng lạ khác - bỏ qua.
            try:
                data = zf.read(media_path)
            except KeyError:
                continue
            anchor_row = int(row_el.text) + 1  # xdr row 0-indexed -> +1 khớp openpyxl
            # Nếu 2 ảnh vô tình khớp cùng 1 anchor_row - giữ ảnh đầu tiên.
            image_map.setdefault(anchor_row, (data, ext))
    return image_map


def _build_col_merge_map(ws, col):
    """Trả về dict {row: value} cho MỌI dòng nằm trong 1 vùng merge DỌC của
    đúng 1 cột `col` (min_col == max_col == col), lấy giá trị từ dòng đầu
    tiên (min_row) của vùng đó."""
    merge_map = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col != col or rng.max_col != col or rng.min_row == rng.max_row:
            continue
        top_val = ws.cell(row=rng.min_row, column=col).value
        for r in range(rng.min_row, rng.max_row + 1):
            merge_map[r] = top_val
    return merge_map


def _resolve_image_block_map(ws, raw_image_map, image_col=11):
    """Ghép {anchor_row: (data, fmt)} thô (theo đúng dòng neo của ảnh) với
    dòng BẮT ĐẦU của khối merge cột `image_col` (K=11, thực tế thường merge
    K:N) chứa dòng neo đó - để khớp đúng với block_start_row bên cột A/B
    (giống hệt logic cũ, chỉ đổi nguồn đọc ảnh thô)."""
    col_merges = [r for r in ws.merged_cells.ranges
                  if r.min_col <= image_col <= r.max_col and r.min_row != r.max_row]
    resolved = {}
    for anchor_row, img in raw_image_map.items():
        block_start_row = anchor_row
        for rng in col_merges:
            if rng.min_row <= anchor_row <= rng.max_row:
                block_start_row = rng.min_row
                break
        resolved.setdefault(block_start_row, img)
    return resolved


def _to_number(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return val
    if isinstance(val, str):
        s = val.strip().replace(',', '').replace(' ', '')
        if not s or s.upper() in ('#N/A', 'N/A', '-'):
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _s(val):
    if val is None:
        return None
    if isinstance(val, str):
        v = val.strip()
        return v or None
    return str(val).strip() or None


def parse_body_kit_excel(path, sheet_name=SHEET_NAME):
    """Trả về (groups, warnings). `groups` là list[dict] với 2 khoá đặc
    biệt: 'parts' (list[dict] các phụ tùng trong nhóm) và 'image' (tuple
    (bytes, format) hoặc None). `warnings` là list[str] các dòng dữ liệu
    bất thường (không đúng 2 dạng "dòng phụ tùng" / "dòng tổng") để người
    dùng soát lại nếu cần - KHÔNG làm dừng quá trình import."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Không tìm thấy sheet '{sheet_name}' trong file. Các sheet có: {wb.sheetnames}")
    ws = wb[sheet_name]

    label_map = _build_col_merge_map(ws, 1)   # Mã loại
    model_map = _build_col_merge_map(ws, 2)   # Model
    raw_image_map = _build_image_row_map_from_zip(path, sheet_name)  # {anchor_row: (bytes, fmt)}
    image_map = _resolve_image_block_map(ws, raw_image_map, image_col=11)  # {block_start_row: (bytes, fmt)}

    groups = []
    warnings = []
    current = None

    for r in range(2, ws.max_row + 1):
        stt = ws.cell(row=r, column=3).value
        part_code = _s(ws.cell(row=r, column=4).value)
        replacement = _s(ws.cell(row=r, column=5).value)
        part_name = _s(ws.cell(row=r, column=6).value)
        honda_price = _to_number(ws.cell(row=r, column=7).value)
        crm_price = _to_number(ws.cell(row=r, column=8).value)
        stock_qty = _to_number(ws.cell(row=r, column=9).value)
        note = _s(ws.cell(row=r, column=10).value)

        # QUAN TRỌNG: kiểm tra "dòng tổng" TRƯỚC khi xét nhãn nhóm (cột A) -
        # dòng tổng có chữ "Tổng" viết thẳng vào cột A (KHÔNG nằm trong
        # vùng merge của nhóm phía trên), nếu xét nhãn trước sẽ bị hiểu
        # nhầm thành 1 nhóm mới tên "Tổng".
        is_total_row = stt is None and part_code is None and (honda_price is not None or crm_price is not None)
        if is_total_row:
            if current is not None:
                current['total_honda_price'] = honda_price
                current['total_crm_price'] = crm_price
            else:
                warnings.append(f'Dòng {r}: gặp dòng tổng nhưng chưa có nhóm nào đang xử lý - đã bỏ qua.')
            continue

        raw_label = ws.cell(row=r, column=1).value
        label = _s(label_map.get(r, raw_label))
        model = _s(model_map.get(r, ws.cell(row=r, column=2).value))

        if not any([label, model, stt, part_code, part_name, honda_price, crm_price]):
            continue  # dòng trắng hoàn toàn -> bỏ qua

        # Bắt đầu 1 nhóm mới khi nhãn (đã resolve merge) đổi khác nhóm
        # đang xử lý - hoặc nhóm đầu tiên của cả sheet.
        if label and (current is None or label != current['group_label']):
            vehicle_family, sub_model = classify_group_label(label)
            current = {
                'group_label': label,
                'model_code': model,
                'vehicle_family': vehicle_family,
                'sub_model': sub_model,
                'year': _extract_year(label),
                'total_honda_price': None,
                'total_crm_price': None,
                'image': image_map.get(r),
                'parts': [],
                '_start_row': r,
            }
            groups.append(current)

        if current is None:
            # Dữ liệu bắt đầu mà chưa từng thấy nhãn nhóm nào - dữ liệu gốc
            # bất thường, ghi nhận cảnh báo và bỏ qua dòng này.
            warnings.append(f'Dòng {r}: chưa xác định được nhóm (thiếu "Mã loại") - đã bỏ qua.')
            continue

        is_part_row = stt is not None and part_code is not None

        if is_part_row:
            current['parts'].append({
                'seq': int(stt) if isinstance(stt, (int, float)) else None,
                'part_code': part_code,
                'replacement_code': replacement,
                'part_name': part_name,
                'honda_price': honda_price,
                'crm_price': crm_price,
                'stock_qty': stock_qty,
                'note': note,
            })
        elif label and stt is None and part_code is None:
            # Đây chính là dòng đầu nhóm (chỉ có Mã loại/Model, chưa có phụ
            # tùng nào) - bình thường, không phải lỗi.
            pass
        else:
            warnings.append(f"Dòng {r}: không khớp dạng dòng phụ tùng lẫn dòng tổng - đã bỏ qua "
                             f"(STT={stt!r}, Mã chi tiết={part_code!r}).")

    # Nhóm không có phụ tùng nào (do lỗi dữ liệu gốc) - vẫn giữ lại (có thể
    # chỉ có ảnh + tên, người dùng tự soát), không tự ý xoá khỏi kết quả.
    for g in groups:
        del g['_start_row']

    return groups, warnings


# ----------------------------------------------------------------------------
# BẢNG TRA "MÃ XE" -> "DÒNG XE / ĐỜI XE" (nhập riêng, xem parse_model_category_
# excel bên dưới) - đây là NGUỒN CHUẨN (do người dùng tự soát/lập), ưu tiên
# CAO HƠN suy luận tự động classify_group_label() ở trên khi 2 nguồn khác
# nhau, vì classify_group_label chỉ đoán từ text tự do của "Mã loại" (thường
# lẫn cả biến thể STD/DX/Magnet/Repsol... vào sub_model, KHÔNG tách theo
# dung tích 110/125 như file tra cứu này).
#
# CẤU TRÚC FILE GỐC (sheet đầu tiên, đã kiểm tra thực tế trên file mẫu):
#   Cột A "Stt"    - số thứ tự BÊN TRONG 1 dòng xe (reset ở mỗi dòng xe mới),
#                    merge dọc theo khối 1 xe/model giống hệt file bảng giá.
#                    RIÊNG dòng ĐẦU 1 nhóm "dòng xe" (vd "AIR BLADE  125"),
#                    cột A chứa THẲNG tên dòng xe đó (không phải số), và cột
#                    B/D của CHÍNH dòng đó luôn trống - đây là dấu hiệu duy
#                    nhất để nhận diện 1 dòng header (không có cột riêng).
#   Cột B:C "Loại xe" (merge ngang B:C) - mô tả model (vd "K27G (Air blade
#                    125) FI DX"), merge dọc xuyên khối 1 model.
#   Cột D "Mã xe"  - mã model ngắn (vd "M3") - merge dọc xuyên khối 1 model.
# LƯU Ý: 1 vài dòng xe bị xuống dòng thành 2 DÒNG HEADER LIÊN TIẾP do copy
# nhầm từ ô có wrap text (vd "WAVE 100" rồi "cc" ở dòng khác) - module này
# GỘP các dòng header liên tiếp (không có dữ liệu model nào chen giữa)
# thành 1 tên dòng xe duy nhất trước khi dùng.
# ----------------------------------------------------------------------------

MODEL_CATEGORY_SHEET_NAME = 'Sheet1'


def _build_col_merge_map_span(ws, col):
    """Giống `_build_col_merge_map` nhưng dùng cho cột nằm TRONG 1 vùng
    merge NGANG nhiều cột (vd B:C) thay vì merge dọc đúng 1 cột - so khớp
    bằng `min_col <= col <= max_col` giống `_resolve_image_block_map`."""
    merge_map = {}
    for rng in ws.merged_cells.ranges:
        if not (rng.min_col <= col <= rng.max_col) or rng.min_row == rng.max_row:
            continue
        top_val = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            merge_map[r] = top_val
    return merge_map


def parse_model_category_excel(path, sheet_name=MODEL_CATEGORY_SHEET_NAME):
    """Đọc file tra "Mã xe" -> "Dòng xe" (vd M3 -> Air Blade 125). Trả về
    (mapping, conflicts, warnings):
        mapping   - dict {model_code: {'vehicle_family', 'sub_model',
                    'raw_header'}}. `vehicle_family` dùng cho menu cấp 1,
                    `sub_model` dùng cho menu cấp 2 (suy luận từ
                    `raw_header` qua classify_group_label() ở trên; nếu
                    không khớp từ khoá dòng xe nào đã biết thì TỰ nó là 1
                    dòng xe cấp 1 riêng - vd "SPACY", "MONKEY" - KHÔNG bị
                    dồn chung vào "Khác" như classify_group_label mặc định,
                    vì ở đây tên dòng xe đã rõ ràng/sạch sẵn, không cần suy
                    luận từ text tự do).
        conflicts - list[str] các Mã xe xuất hiện dưới >1 dòng xe khác nhau
                    trong cùng file (dữ liệu gốc mâu thuẫn - vẫn dùng lần
                    khớp SAU CÙNG, người dùng tự soát lại nếu cần).
        warnings  - list[str] các dòng bất thường khác (mô tả không có mã,
                    mã xuất hiện trước khi có dòng xe nào...).
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Không tìm thấy sheet '{sheet_name}' trong file. Các sheet có: {wb.sheetnames}")
    ws = wb[sheet_name]

    label_map = _build_col_merge_map_span(ws, 2)   # Loại xe (B:C)
    code_map = _build_col_merge_map(ws, 4)          # Mã xe (D)

    mapping = {}
    conflicts = []
    warnings = []
    current_header = None
    pending_header_parts = []

    # Gom theo KHỐI (mọi dòng liên tiếp có cùng nhãn đã resolve-merge) rồi
    # mới quyết định mã/cảnh báo cho CẢ khối - vì vùng merge của cột "Mã xe"
    # (D) đôi khi HẸP HƠN vùng merge của cột "Loại xe" (B:C) trong cùng 1
    # khối (vd nhãn merge dòng 37-39 nhưng mã chỉ có ở dòng 38) - nếu xét
    # từng dòng riêng lẻ sẽ báo nhầm "có mô tả nhưng thiếu mã" cho dòng 37/39
    # dù cả khối THỰC RA có mã (ở dòng 38).
    block_rows = []  # [(row, code_or_None, label_or_None), ...] của khối đang gom
    block_label = None

    def _flush_block():
        nonlocal block_rows, block_label
        if not block_rows:
            return
        code = next((c for _, c, _ in block_rows if c), None)
        first_row = block_rows[0][0]
        if not code:
            if block_label:
                warnings.append(f"Dòng {first_row}: có mô tả '{block_label}' nhưng không có Mã xe - đã bỏ qua.")
            block_rows = []
            return
        if current_header is None:
            warnings.append(f"Dòng {first_row}: Mã xe '{code}' xuất hiện trước khi có dòng xe nào - đã bỏ qua.")
            block_rows = []
            return

        prior = mapping.get(code)
        if prior and prior['raw_header'] != current_header:
            conflicts.append(
                f"Mã xe {code}: xuất hiện ở cả '{prior['raw_header']}' và '{current_header}' - "
                f"đang dùng '{current_header}' (lần khớp sau cùng)."
            )

        family, sub_model = classify_group_label(current_header)
        if family == 'Khác':
            # Không khớp từ khoá dòng xe nào đã biết trong _FAMILY_PATTERNS
            # (vd "SPACY", "MONKEY", "SUPER CUB") - nhưng đây là tên dòng xe
            # ĐÃ SẠCH (không phải text tự do lẫn biến thể), nên coi CHÍNH nó
            # là 1 dòng xe cấp 1 riêng thay vì dồn vào "Khác" chung chung.
            family = current_header
            sub_model = current_header

        mapping[code] = {
            'vehicle_family': family,
            'sub_model': sub_model,
            'raw_header': current_header,
        }
        block_rows = []

    for r in range(2, ws.max_row + 1):
        raw_a = ws.cell(row=r, column=1).value
        raw_b = ws.cell(row=r, column=2).value
        raw_d = ws.cell(row=r, column=4).value

        is_header_row = isinstance(raw_a, str) and raw_a.strip() and raw_b is None and raw_d is None
        if is_header_row:
            _flush_block()
            block_label = None
            pending_header_parts.append(raw_a.strip())
            continue

        code = _s(code_map.get(r, raw_d))
        label = _s(label_map.get(r, raw_b))
        if not code and not label:
            continue  # dòng trắng hoàn toàn giữa 2 khối - bỏ qua

        if pending_header_parts:
            current_header = re.sub(r'\s+', ' ', ' '.join(pending_header_parts)).strip()
            pending_header_parts = []

        if label != block_label:
            _flush_block()
            block_label = label
        block_rows.append((r, code, label))

    _flush_block()

    return mapping, conflicts, warnings


if __name__ == '__main__':
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else \
        '/mnt/user-data/uploads/BA_NG_GIA__BO___A_O_XE_HONDA.xlsx'
    groups, warnings = parse_body_kit_excel(path)
    total_parts = sum(len(g['parts']) for g in groups)
    with_image = sum(1 for g in groups if g['image'])
    print(f'Đọc OK: {len(groups)} nhóm, {total_parts} dòng phụ tùng, {with_image} nhóm có ảnh, '
          f'{len(warnings)} cảnh báo.')
    from collections import Counter
    fam_counts = Counter(g['vehicle_family'] for g in groups)
    print('\n--- Số nhóm theo dòng xe ---')
    for fam, cnt in fam_counts.most_common():
        print(f'  {fam}: {cnt}')

    print('\n--- 3 nhóm mẫu đầu ---')
    for g in groups[:3]:
        print(g['group_label'], '|', g['vehicle_family'], '/', g['sub_model'], '| model:', g['model_code'],
              '| parts:', len(g['parts']), '| image:', bool(g['image']))
        for p in g['parts'][:2]:
            print('   ', p)
    if warnings:
        print('\n--- 10 cảnh báo đầu ---')
        for w in warnings[:10]:
            print(w)